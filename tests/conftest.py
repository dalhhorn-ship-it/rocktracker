import datetime as dt
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import rock_model  # noqa: E402
SCHEMA_PATH = ROOT / "schema" / "workbook.schema.json"


@pytest.fixture(scope="session")
def schema():
    return json.loads(SCHEMA_PATH.read_text())


@pytest.fixture(scope="session")
def sheets(schema):
    return {s["name"]: s for s in schema["sheets"]}


@pytest.fixture(scope="session")
def workbook(schema):
    path = ROOT / schema["workbook"]
    if not path.exists():
        pytest.skip(f"{schema['workbook']} not found; run scripts/build_workbook.py")
    return load_workbook(path, data_only=True)


@pytest.fixture(scope="session")
def sample_today():
    """The frozen 'today' of the sample (data/sample.json)."""
    return dt.date.fromisoformat(json.loads((ROOT / "data" / "sample.json").read_text())["today"])


@pytest.fixture(scope="session")
def rows(workbook, schema, sample_today):
    """Sheet name to list of row dicts keyed by header.

    Table sheets (history, lookup) are read from the workbook. The five consolidated sheets
    hold formulas that nothing outside Excel can calculate, so their rows come from the
    Python mirror (scripts/rock_model.py) fed with the typed cells of every rock tab in the
    workbook, with today frozen to the sample date."""
    out = {}
    for sheet in schema["sheets"]:
        if sheet["kind"] == "consolidated" or sheet["name"] not in workbook.sheetnames:
            continue
        values = list(workbook[sheet["name"]].iter_rows(values_only=True))
        headers = list(values[0]) if values else []
        out[sheet["name"]] = [
            dict(zip(headers, r)) for r in values[1:] if any(v is not None for v in r)
        ]
    rocks = rock_model.rocks_from_workbook(ROOT / schema["workbook"], schema)
    steercos = [r["Date"] for r in out["lkp_Steercos"]]
    out.update(rock_model.build_model(rocks, sample_today, steercos, schema))
    return out


def split_ref(ref):
    sheet, column = ref.split(".", 1)
    return sheet, column
