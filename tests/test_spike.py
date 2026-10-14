"""T05: the consolidation spike file and the dynamic array helper."""
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import build_spike  # noqa: E402
from build_workbook import to_file_formula  # noqa: E402


@pytest.fixture(scope="module")
def spike(tmp_path_factory):
    return build_spike.build(tmp_path_factory.mktemp("spike") / "spike.xlsx")


def _read(path, name):
    with zipfile.ZipFile(path) as z:
        return z.read(name).decode()


def test_prefixes():
    out = to_file_formula('=LET(d,VSTACK(S!A1),FILTER(d,d<>"d"))', "d")
    assert out == '=_xlfn.LET(_xlpm.d,_xlfn.VSTACK(S!A1),_xlfn._xlws.FILTER(_xlpm.d,_xlpm.d<>"d"))'


def test_sheet_order(spike):
    names = load_workbook(spike).sheetnames
    assert names.index("RockTemplate") < names.index("R01") < names.index("R02") < names.index("RocksEnd")
    assert {"Epics", "Epics_C2", "RockTabs"} <= set(names)


def test_anchor_is_dynamic_array(spike):
    names = load_workbook(spike).sheetnames
    for sheet in ("Epics", "Epics_C2"):
        xml = _read(spike, f"xl/worksheets/sheet{names.index(sheet) + 1}.xml")
        m = re.search(r'<c r="A2" cm="1"><f t="array" ref="A2">(.*?)</f>', xml)
        assert m, f"{sheet} A2 lacks cm=1 and array formula"
        assert "_xlfn.VSTACK" in m.group(1) and "_xlpm.d" in m.group(1)
    assert "RockTemplate:RocksEnd!A5:D9" in _read(spike, f"xl/worksheets/sheet{names.index('Epics') + 1}.xml")
    assert "_xlfn.REDUCE" in _read(spike, f"xl/worksheets/sheet{names.index('Epics_C2') + 1}.xml")
    assert "XLDAPR" in _read(spike, "xl/metadata.xml")
    assert "sheetMetadata" in _read(spike, "[Content_Types].xml")
    assert "metadata.xml" in _read(spike, "xl/_rels/workbook.xml.rels")


def test_schema_formula_used(spike, schema):
    cf = schema["rock_tab"]["consolidation_formula"]
    assert cf["span"] == "RockTemplate:RocksEnd"
    assert build_spike.plan_a_formula(schema).startswith("=LET(d,VSTACK(RockTemplate:RocksEnd!")


@pytest.mark.skipif(not shutil.which("soffice"), reason="soffice not installed")
def test_libreoffice_opens(spike, tmp_path):
    r = subprocess.run(["soffice", "--headless", "--convert-to", "csv", "--outdir", str(tmp_path), str(spike)],
                       capture_output=True, text=True, timeout=120)
    assert r.returncode == 0 and list(tmp_path.glob("*.csv")), r.stderr
