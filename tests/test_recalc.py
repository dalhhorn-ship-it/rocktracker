"""Layer 7 (T09, TD11): recalculate the sample workbook headless in LibreOffice and compare every
rock tab calculated cell with the Python mirror (scripts/rock_model.py).

LibreOffice cannot freeze TODAY(), so the model gets the same date LibreOffice uses (the
machine date). The six consolidated sheets use VSTACK, which LibreOffice 24.2 lacks, so they
are skipped explicitly (the rock tab formulas that feed them are what we compare).

Also recomputes the sample snapshot rows with the model (TD11)."""
import datetime as dt
import json
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import rock_model as rm
import snapshot as sn

ROOT = Path(__file__).resolve().parent.parent
SAMPLE = ROOT / "data" / "sample.json"


# ---------------------------------------------------------------- helpers
def norm(v):
    if v is None or v == "":
        return ""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return float(v)
    return v


def same(a, b):
    a, b = norm(a), norm(b)
    if isinstance(a, float) and isinstance(b, float):
        return abs(a - b) < 1e-9
    return a == b


@pytest.fixture(scope="module")
def recalculated(schema, tmp_path_factory):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        pytest.skip("soffice (LibreOffice) not installed; recalculation check skipped")
    src = ROOT / schema["workbook"]
    if not src.exists():
        pytest.skip(f"{schema['workbook']} not found; run scripts/build_workbook.py")
    out = tmp_path_factory.mktemp("recalc")
    profile = tmp_path_factory.mktemp("lo_profile")
    today = dt.date.today()
    cmd = [soffice, f"-env:UserInstallation=file://{profile}", "--headless",
           "--convert-to", "xlsx", "--outdir", str(out), str(src)]
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=90)
    target = out / src.name
    assert target.exists(), f"LibreOffice conversion failed: {res.stdout} {res.stderr}"
    return {"wb": load_workbook(target, data_only=True), "today": today, "src": src}


@pytest.fixture(scope="module")
def steercos(recalculated):
    ws = recalculated["wb"]["lkp_Steercos"]
    return [r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]


@pytest.fixture(scope="module")
def tabs(recalculated, schema, steercos):
    """List of (sheet name, rock dict from typed cells, model result for that rock alone)."""
    rocks = rm.rocks_from_workbook(recalculated["src"], schema)
    names = recalculated["wb"].sheetnames
    rt = schema["rock_tab"]
    tab_names = names[names.index(rt["template_sheet"]) + 1:names.index(rt["end_sheet"])]
    assert len(tab_names) == len(rocks) and rocks, "no rock tabs found"
    return [(name, rock, rm.build_model([rock], recalculated["today"], steercos, schema))
            for name, rock in zip(tab_names, rocks)]


# ---------------------------------------------------------------- rock tab cells
def test_consolidated_sheets_skipped_explicitly(schema):
    """Documented skip: consolidated sheets need VSTACK, FILTER, CHOOSECOLS and LET with
    SEQUENCE, which LibreOffice 24.2 lacks (TD recalc item). The rock tab formulas that feed
    them are compared instead, including the Teams columns the keep_zero_columns guard reads."""
    skipped = [s["name"] for s in schema["sheets"] if s["kind"] == "consolidated"]
    assert set(skipped) == {"Rocks", "Milestones", "Risks", "CallOuts", "Decisions", "Teams"}


def test_strip_row_matches_model(recalculated, schema, tabs):
    """Hidden strip row 2 (all Rocks columns, incl. CurrentMCQuarter, NextMCQuarter, section 7
    totals, FillPct, signals) against the model's Rocks row."""
    rt = schema["rock_tab"]
    cols = [c["name"] for c in next(s for s in schema["sheets"] if s["name"] == "Rocks")["columns"]]
    bad = []
    for name, rock, model in tabs:
        ws = recalculated["wb"][name]
        cells = [ws.cell(row=rt["strip_row"], column=i + 1).value for i in range(len(cols))]
        if not model["Rocks"]:
            assert all(norm(v) == "" for v in cells), f"{name}: strip not blank without RockID"
            continue
        exp = model["Rocks"][0]
        bad += [f"{name}.{c}: excel={v!r} model={exp[c]!r}"
                for c, v in zip(cols, cells) if not same(v, exp[c])]
    assert not bad, "strip mismatches:\n" + "\n".join(bad)


def test_calculated_fields_match_model(recalculated, schema, tabs, steercos):
    """Every formula field of the rock tab sections 1, 2 and 7 (OpenRequests, Request1 to 5,
    signals, totals) against rock_model.tab_fields."""
    rt = schema["rock_tab"]
    fields = [f for sec in rt["sections"] if sec["kind"] == "fields" for f in sec["fields"]
              if f.get("formula")]
    names = {f["name"] for f in fields}
    assert {"OpenRequests", "Request1", "Request5", "UpdateAgeDays", "DataSignal"} <= names
    expected = rm.tab_fields([rock for _, rock, _ in tabs], recalculated["today"], steercos)
    bad, checked = [], 0
    for (name, rock, model), exp in zip(tabs, expected):
        ws = recalculated["wb"][name]
        for f in fields:
            got = ws[f"{rt['value_column']}{f['row']}"].value
            checked += 1
            if f["name"] not in exp:
                bad.append(f"{name}.{f['name']}: not in tab_fields")
            elif not same(got, exp[f["name"]]):
                bad.append(f"{name}.{f['name']}: excel={got!r} model={exp[f['name']]!r}")
    assert checked and not bad, "tab field mismatches:\n" + "\n".join(bad)


@pytest.mark.parametrize("block", ["Milestones", "Risks", "CallOuts", "Decisions", "Teams"])
def test_block_calculated_columns_match_model(recalculated, schema, tabs, block):
    """Calculated columns of each block (Milestones: Status, SlipWeeks, TimelineDate, MCQuarter,
    InOutlook, Visible, Overdue; Decisions: DecisionState, RequestNo; Teams: FillPct) per row.
    Only the rock tab side: the consolidated zero wrap (keep_zero_columns) is VSTACK based and
    is checked on the mirror and by the anchor formula test."""
    rt = schema["rock_tab"]
    sec = next(s for s in rt["sections"] if s.get("block") == block)
    sheet = next(s for s in schema["sheets"] if s["name"] == block)
    calc = [c for c in sheet["columns"] if c.get("formula")]
    key = calc[0]["name"]
    assert key.endswith("Key")
    checked, bad = 0, []
    for name, rock, model in tabs:
        ws = recalculated["wb"][name]
        headers = {str(c.value).strip(): c.column for c in ws[sec["column_header_row"]] if c.value}
        by_key = {r[key]: r for r in model[block]}
        keyed = 0
        for r in range(sec["first_row"], sec["first_row"] + sec["capacity"]):
            kv = ws.cell(row=r, column=headers[next(c["label"] for c in calc if c["name"] == key)]).value
            if norm(kv) == "":
                continue
            keyed += 1
            exp = by_key.get(kv)
            if exp is None:
                bad.append(f"{name} row {r}: key {kv!r} not in model")
                continue
            for c in calc:
                got = ws.cell(row=r, column=headers[c["label"]]).value
                checked += 1
                want = exp[c["name"]]
                if c["name"] == "FillPct" and want == 0:
                    want = 0  # keep zero column: real 0 stays 0
                if not same(got, want):
                    bad.append(f"{name} row {r} {c['name']}: excel={got!r} model={want!r}")
        if keyed != len(model[block]):
            bad.append(f"{name}: {keyed} keyed rows in Excel, {len(model[block])} in model")
    assert not bad, f"{block} mismatches:\n" + "\n".join(bad)
    assert checked > 0, f"no {block} cells compared"


def test_milestone_cases_are_covered(tabs):
    """The sample must exercise the 2.1.0 columns this check compares."""
    ms = [m for _, _, model in tabs for m in model["Milestones"]]
    assert {m["Visible"] for m in ms} == {"Yes", "No"}
    assert {m["InOutlook"] for m in ms} >= {"Yes", "No"}
    assert len({m["MCQuarter"] for m in ms}) >= 3
    teams = [t for _, _, model in tabs for t in model["Teams"]]
    assert any(t["Budgeted"] == 0 for t in teams) and teams


# ---------------------------------------------------------------- snapshots (TD11)
def _snapshot_rows(sample, schema, sheet):
    cols = [c["name"] for c in next(s for s in schema["sheets"] if s["name"] == sheet)["columns"]]
    return [dict(zip(cols, r)) for r in sample["snapshots"][sheet]]


def test_latest_snapshot_recomputed_by_model(schema):
    """snapshot.py applied to the model output has exactly the history sheet columns, and the
    newest stored snapshot agrees on the fields that do not move with the week (identity,
    positions, high risks). The milestone counts of the stored sample snapshots are static
    synthetic history (sample data is from 12 Oct, the model runs on the sample today of 14 Oct),
    so they are checked for internal consistency below, not recomputed (TECH_DEBT)."""
    sample = json.loads(SAMPLE.read_text())
    steercos = [r[0] for r in sample["lookups"]["lkp_Steercos"]]
    stored_r = _snapshot_rows(sample, schema, "RockSnapshots")
    stored_m = _snapshot_rows(sample, schema, "MilestoneSnapshots")
    today = dt.date.fromisoformat(sample["today"])
    model = rm.build_model(sample["rocks"], today, steercos, schema)
    new_r, new_m = sn.snapshot_rows(model["Rocks"], model["Milestones"], today)
    cols = {s["name"]: [c["name"] for c in s["columns"]] for s in schema["sheets"]}
    assert new_r and new_m
    assert all(list(r) == cols["RockSnapshots"] for r in new_r)
    assert all(list(r) == cols["MilestoneSnapshots"] for r in new_m)
    latest = max(r["SnapshotDate"] for r in stored_r)
    stable = ("RockID", "Rock", "GPM", "PositionsBudgeted", "PositionsActual", "OpenHighRisks")
    by_rock = {r["RockID"]: r for r in new_r}
    bad = []
    for r in (r for r in stored_r if r["SnapshotDate"] == latest):
        exp = by_rock.get(r["RockID"])
        if exp is None:
            bad.append(f"{r['SnapshotID']} missing from recompute")
            continue
        bad += [f"{r['SnapshotID']}.{c}: stored={r[c]!r} model={exp[c]!r}"
                for c in stable if not same(r[c], exp[c])]
    assert not bad, "\n".join(bad)


def test_earlier_snapshots_internally_consistent(schema):
    """Earlier weeks cannot be recomputed from today's rock tabs (the data moved on), so check
    that each stored milestone status follows the status rule from its stored dates and commitment."""
    sample = json.loads(SAMPLE.read_text())
    bad = []
    for r in _snapshot_rows(sample, schema, "MilestoneSnapshots"):
        assert r["SnapshotID"] == f"{r['SnapshotDate']}|{r['MilestoneKey']}"
        assert r["MilestoneKey"].startswith(r["RockID"] + "-M")
        planned, forecast = rm.to_date(r["PlannedDate"]), rm.to_date(r["ForecastDate"])
        status, commit = r["Status"], r["Commitment"]
        if status in ("Completed", "On hold"):
            continue
        if commit != "Committed":
            ok = status == "Not committed"
        elif rm.n(planned) > 0 and rm.n(forecast) > rm.n(planned):
            ok = status == "Slipped"
        else:
            ok = status in ("On plan", "On watch")
        if not ok:
            bad.append(f"{r['SnapshotID']}: {commit} {planned} {forecast} -> {status}")
    assert not bad, "\n".join(bad)
