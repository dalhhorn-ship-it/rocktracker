"""Layer 5: the deployment workbook is ready for the shared drive, and migration keeps data."""
import sys

import pytest
from openpyxl import load_workbook

from conftest import ROOT

sys.path.insert(0, str(ROOT / "scripts"))
import build_workbook  # noqa: E402


@pytest.fixture(scope="module")
def deploy_wb(tmp_path_factory):
    out = tmp_path_factory.mktemp("deploy") / "deploy.xlsx"
    build_workbook.main(["--deploy", "--out", str(out)])
    return load_workbook(out)


def validations(ws):
    return {str(r): dv for dv in ws.data_validations.dataValidation for r in dv.sqref.ranges}


def test_about_sheet_first(deploy_wb, schema):
    assert deploy_wb.sheetnames[0] == "About"
    listed = {deploy_wb["About"].cell(row=r, column=1).value for r in range(1, 40)}
    assert {s["name"] for s in schema["sheets"]} <= listed


def test_every_sheet_is_a_table_with_frozen_header(deploy_wb, schema):
    for s in schema["sheets"]:
        ws = deploy_wb[s["name"]]
        assert f"tbl_{s['name']}" in ws.tables
        assert ws.freeze_panes == "A2"


def test_lookups_seeded_and_no_sample_rows(deploy_wb, schema):
    for s in schema["sheets"]:
        data_rows = [r for r in deploy_wb[s["name"]].iter_rows(min_row=2, values_only=True) if any(r)]
        assert len(data_rows) == len(s.get("seed", [])), f"{s['name']} should ship with seed rows only"


def test_dropdowns_on_lookup_and_reference_columns(deploy_wb, schema):
    for s in schema["sheets"]:
        dvs = validations(deploy_wb[s["name"]])
        for i, c in enumerate(s["columns"]):
            ref = c.get("lookup") or c.get("references")
            if not ref:
                continue
            letter = chr(ord("A") + i)
            match = [dv for rng, dv in dvs.items() if rng.startswith(f"{letter}2:")]
            assert match and match[0].type == "list", f"{s['name']}.{c['name']} has no dropdown"
            assert build_workbook.list_name(ref) in deploy_wb.defined_names


def test_migration_keeps_rows(tmp_path, workbook, schema):
    out = tmp_path / "migrated.xlsx"
    build_workbook.main(["--deploy", "--from", str(ROOT / schema["workbook"]), "--out", str(out)])
    migrated = load_workbook(out)
    for s in schema["sheets"]:
        before = sum(1 for r in workbook[s["name"]].iter_rows(min_row=2, values_only=True) if any(r))
        after = sum(1 for r in migrated[s["name"]].iter_rows(min_row=2, values_only=True) if any(r))
        assert before == after, f"{s['name']} lost rows in migration"
