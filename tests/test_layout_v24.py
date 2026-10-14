"""Schema 2.4.0 (change request 3): profile first, value chain seed, plain explanations of grey cells,
migration of a 2.3 file to the new rows."""
import json
import re
import sys

import pytest
from openpyxl import load_workbook

from conftest import ROOT

sys.path.insert(0, str(ROOT / "scripts"))
import build_workbook  # noqa: E402

OLD = json.loads((ROOT / "schema" / "history" / "workbook.schema.v2.3.0.json").read_text())
VALUE_CHAINS = ["B2C", "B2B", "Trade", "Assets", "Enterprise"]


def sections(schema):
    return sorted(schema["rock_tab"]["sections"], key=lambda s: s["header_row"])


def span(sec):
    if sec["kind"] == "fields":
        return sec["header_row"], max(f["row"] for f in sec["fields"])
    return sec["header_row"], sec["first_row"] + sec["capacity"] - 1


# ------------------------------------------------------------------ layout
def test_profile_is_the_first_section_directly_under_the_title_strip_and_guide(schema):
    rt = schema["rock_tab"]
    first = sections(schema)[0]
    assert first["kind"] == "fields" and "Rock profile" in first["title"]
    assert {f["name"] for f in first["fields"]} >= {"RockID", "RockName", "ValueChain", "Lead", "GPM", "Active"}
    guide_row = int(re.sub(r"\D", "", rt["guide_cell"]))
    assert first["header_row"] == max(rt["strip_row"], guide_row) + 2  # one blank row after the guide
    assert [f["row"] for f in first["fields"]] == list(range(first["header_row"] + 1, first["header_row"] + 1 + len(first["fields"])))


def test_section_order_ids_and_one_blank_row_between_sections(schema):
    secs = sections(schema)
    names = [s.get("block") or re.sub(r"^\d+\s+", "", s["title"]).split(" ")[0] for s in secs]
    assert names == ["Rock", "Update", "Signals", "Milestones", "Risks", "CallOuts", "Decisions", "Teams"]
    assert [s["id"] for s in secs] == list(range(1, 9))
    for s in secs:
        assert s["title"].startswith(f"{s['id']}  ")
    for a, b in zip(secs, secs[1:]):
        assert b["header_row"] == span(a)[1] + 2, f"section {a['id']} and {b['id']} are not one blank row apart"
    spans = [span(s) for s in secs]
    assert all(x[1] < y[0] for x, y in zip(spans, spans[1:])), "overlap"


def test_text_references_to_sections_follow_the_new_numbers(schema):
    text = json.dumps(schema)
    assert "section 7 Rock profile" not in text and "section 7 profile" not in text
    assert "section 1 Rock profile" in text
    by_block = {s["block"]: s["id"] for s in schema["rock_tab"]["sections"] if s["kind"] == "block"}
    for s in schema["sheets"]:
        m = re.search(r"stacked from section (\d+)", s.get("purpose", ""))
        if m:
            assert int(m.group(1)) == by_block[s["name"]], s["name"]


def test_headers_the_portal_reads_are_unchanged_from_2_3(schema):
    new = {s["name"]: [c["name"] for c in s["columns"]] for s in schema["sheets"]}
    old = {s["name"]: [c["name"] for c in s["columns"]] for s in OLD["sheets"]}
    assert new == old


def test_built_tabs_follow_the_schema_rows(workbook, schema):
    rt = schema["rock_tab"]
    ws = workbook[rt["template_sheet"]]
    for sec in rt["sections"]:
        assert ws[f"{rt['label_column']}{sec['header_row']}"].value == sec["title"]
        if sec["kind"] == "fields":
            for f in sec["fields"]:
                assert ws[f"{rt['label_column']}{f['row']}"].value == f["label"]
    assert ws["B5"].value.startswith("1  Rock profile") and ws["B6"].value.startswith("Rock ID")


# ------------------------------------------------------------------ value chains
def test_value_chain_seed_is_exactly_the_five_in_order(schema, sheets):
    assert [r[0] for r in sheets["lkp_ValueChains"]["seed"]] == VALUE_CHAINS


def test_sample_rocks_use_only_the_seeded_value_chains():
    sample = json.loads((ROOT / "data" / "sample.json").read_text())
    assert sample["lookups"]["lkp_ValueChains"] == []
    assert {r["ValueChain"] for r in sample["rocks"]} <= set(VALUE_CHAINS)


def test_value_chain_lookup_in_both_workbooks(workbook, tmp_path):
    def chains(wb):
        return [r[0] for r in wb["lkp_ValueChains"].iter_rows(min_row=2, values_only=True) if r[0]]

    assert chains(workbook) == VALUE_CHAINS
    out = tmp_path / "d.xlsx"
    build_workbook.main(["--deploy", "--out", str(out)])
    assert chains(load_workbook(out)) == VALUE_CHAINS


# ------------------------------------------------------------------ plain explanations
def calculated(schema):
    for sec in schema["rock_tab"]["sections"]:
        for f in sec.get("fields", []):
            if f.get("formula"):
                yield f"field {f['name']}", f
    for s in schema["sheets"]:
        for c in s["columns"]:
            if c.get("formula"):
                yield f"{s['name']}.{c['name']}", c


def test_every_calculated_field_and_column_has_a_plain_sentence(schema):
    n = 0
    for where, c in calculated(schema):
        n += 1
        p = c.get("plain")
        assert p, f"{where} has no plain explanation"
        assert p.endswith(".") and len(p) <= 260, where
        assert not re.search(r"[=\[\]{}]|\b[A-Z]+\(", p), f"{where}: plain text looks like a formula"
        assert not re.search(r"\.\s+[A-Z]", p), f"{where}: plain text is more than one sentence"
    assert n >= 70


def test_plain_text_of_the_two_given_examples(sheets):
    risks = {c["name"]: c for c in sheets["Risks"]["columns"]}
    assert risks["High"]["plain"] == "Yes when likelihood times impact scores 6 or more (Low 1, Medium 2, High 3), so High x Medium, Medium x High or High x High."
    assert risks["NeedsAction"]["plain"] == "Yes when the risk is High, not closed, and the key action or the action owner is empty."


@pytest.fixture(scope="module")
def tabs(workbook, schema):
    rt = schema["rock_tab"]
    names = workbook.sheetnames
    return [workbook[n] for n in names[names.index(rt["template_sheet"]):names.index(rt["end_sheet"])]]


def test_plain_text_sits_in_column_d_next_to_every_grey_field(workbook, schema):
    rt = schema["rock_tab"]
    wb = load_workbook(ROOT / schema["workbook"])
    names = wb.sheetnames
    for name in names[names.index(rt["template_sheet"]):names.index(rt["end_sheet"])]:
        ws = wb[name]
        for sec in rt["sections"]:
            for f in sec.get("fields", []):
                d = ws[f"D{f['row']}"]
                if f.get("formula"):
                    assert d.value == f["plain"], (name, f["name"])
                    assert d.font.i and d.font.sz == 9 and d.font.color.rgb.endswith("72777D"), (name, f["name"])
                else:
                    assert d.value is None, (name, f["name"])


def test_grey_column_headers_carry_a_note_with_the_plain_text(schema):
    rt = schema["rock_tab"]
    wb = load_workbook(ROOT / schema["workbook"])
    names = wb.sheetnames
    sheets = {s["name"]: s for s in schema["sheets"]}
    for name in names[names.index(rt["template_sheet"]):names.index(rt["end_sheet"])]:
        ws = wb[name]
        for sec in rt["sections"]:
            if sec["kind"] != "block":
                continue
            for i, c in enumerate(sheets[sec["block"]]["columns"], start=1):
                cell = ws.cell(row=sec["column_header_row"], column=i)
                if c.get("formula"):
                    assert cell.comment is not None and c["plain"] in cell.comment.text, (name, sec["block"], c["name"])
                else:
                    assert cell.comment is None, (name, sec["block"], c["name"])


def test_about_has_the_how_grey_cells_are_calculated_table(schema, tmp_path):
    out = tmp_path / "d.xlsx"
    build_workbook.main(["--deploy", "--out", str(out)])
    for about in (load_workbook(out)["About"], load_workbook(ROOT / schema["workbook"])["About"]):
        rows = [[c.value for c in r] for r in about.iter_rows()]
        titles = [i for i, r in enumerate(rows) if r[0] == "How the grey cells are calculated"]
        assert len(titles) == 1
        head = rows[titles[0] + 1]
        assert head[:2] == ["Field", "Plain explanation"]
        listed = []
        for r in rows[titles[0] + 2:]:
            if not r[0]:
                break
            listed.append(tuple(r[:3]))
        expected = build_workbook.grey_cells(schema)
        assert listed == expected and len(listed) > 40
        table = {(a, b): c for a, b, c in listed}
        # every calculated field of the rock tab is in it
        for sec in schema["rock_tab"]["sections"]:
            for f in sec.get("fields", []):
                if f.get("formula"):
                    assert (f["label"], f["plain"]) in table


# ------------------------------------------------------------------ migration 2.3 to 2.4
def test_migration_from_a_2_3_file_moves_rows_and_keeps_every_value(tmp_path, schema):
    sample = json.loads((ROOT / "data" / "sample.json").read_text())
    rows = {s["name"]: [list(r) for r in s.get("seed", [])] for s in OLD["sheets"] if s["kind"] != "consolidated"}
    rows["lkp_People"] = [["Alex Example", "Lead"]] + rows["lkp_People"]
    rows["lkp_ValueChains"] = [["Checkout"], ["Fulfilment"]]  # the old live values stay in a migrated file
    old_path = tmp_path / "v23.xlsx"
    build_workbook.build(OLD, {"rocks": sample["rocks"], "rows": rows}, deploy=False).save(old_path)
    wb = load_workbook(old_path)
    assert wb["RockTemplate"]["B105"].value.startswith("7  Rock profile")  # really the old layout

    out = tmp_path / "v24.xlsx"
    build_workbook.main(["--deploy", "--from", str(old_path), "--out", str(out)])
    new = load_workbook(out)
    assert new["RockTemplate"]["B5"].value.startswith("1  Rock profile")
    before = build_workbook.read_rock_tabs(old_path, OLD)
    after = build_workbook.read_rock_tabs(out, schema)
    assert before and before == after
    rid = next(f["row"] for f in schema["rock_tab"]["sections"][0]["fields"] if f["name"] == "RockID")
    assert new[after[0]["RockTab"]][f"C{rid}"].value == after[0]["RockID"]
    assert [r[0] for r in new["lkp_ValueChains"].iter_rows(min_row=2, values_only=True) if r[0]] == ["Checkout", "Fulfilment"]
    assert schema["migration"]["v2_3_to_v2_4"]
