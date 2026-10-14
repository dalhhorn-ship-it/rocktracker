"""T07: the generator output for rock tabs, consolidated sheets, names and protection."""
import re
import shutil
import subprocess
import sys
import zipfile

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from conftest import ROOT

sys.path.insert(0, str(ROOT / "scripts"))
import build_workbook  # noqa: E402

SNAPSHOT_TS = (ROOT / "scripts" / "office" / "snapshot.ts").read_text()


@pytest.fixture(scope="module")
def wb():
    return load_workbook(ROOT / "data" / "workbook.xlsx")


@pytest.fixture(scope="module")
def tabs(wb, schema):
    rt = schema["rock_tab"]
    names = wb.sheetnames
    return names[names.index(rt["template_sheet"]):names.index(rt["end_sheet"])]


def fixed(cell):
    return cell.value.text if isinstance(cell.value, ArrayFormula) else cell.value


def test_no_placeholder_left_in_any_formula(wb, tabs, schema):
    bad = []
    for name in tabs + [s["name"] for s in schema["sheets"] if s["kind"] == "consolidated"]:
        for row in wb[name].iter_rows():
            for c in row:
                v = fixed(c)
                if isinstance(v, str) and v.startswith("=") and re.search(r"[{}\[\]]", v):
                    bad.append(f"{name}!{c.coordinate}")
    assert not bad, bad[:10]
    for cf in (wb[t].conditional_formatting for t in tabs):
        for rng in cf:
            for rule in rng.rules:
                assert not re.search(r"[{}\[\]]", rule.formula[0]), rule.formula


def test_strip_row_hidden_and_in_rocks_column_order(wb, schema, tabs):
    rt = schema["rock_tab"]
    rocks = next(s for s in schema["sheets"] if s["name"] == "Rocks")
    for t in tabs:
        ws = wb[t]
        assert ws.row_dimensions[rt["strip_row"]].hidden
        strip = [c.value for c in ws[rt["strip_row"]]]
        assert len(strip) == len(rocks["columns"]) and all(str(v).startswith("=") for v in strip)


def test_calculated_cells_grey_locked_and_entry_cells_unlocked(wb, schema, tabs):
    rt = schema["rock_tab"]
    sheets = {s["name"]: s for s in schema["sheets"]}
    ws = wb[rt["template_sheet"]]
    assert ws.protection.sheet and ws.protection.insertRows and ws.protection.deleteRows
    assert ws.protection.insertColumns and ws.protection.deleteColumns
    for sec in rt["sections"]:
        if sec["kind"] == "fields":
            cells = [(ws[f"{rt['value_column']}{f['row']}"], f) for f in sec["fields"]]
        else:
            r = sec["first_row"] + 2
            cells = [(ws.cell(row=r, column=i), c) for i, c in enumerate(sheets[sec["block"]]["columns"], 1)]
        for cell, c in cells:
            if c.get("formula"):
                assert cell.protection.locked and cell.fill.fgColor.rgb.endswith(build_workbook.CALC_GREY), c["name"]
                assert str(cell.value).startswith("=")
            else:
                assert not cell.protection.locked, c["name"]


def field_cells(schema):
    rt = schema["rock_tab"]
    return {f["name"]: f"{rt['value_column']}{f['row']}" for sec in rt["sections"] for f in sec.get("fields", [])}


def block(schema, name):
    return next(sec for sec in schema["rock_tab"]["sections"] if sec.get("block") == name)


def test_pink_rules_present(wb, schema, sheets):
    ws = wb[schema["rock_tab"]["template_sheet"]]
    rules = {str(r.sqref): [x.formula[0] for x in r.rules] for r in ws.conditional_formatting}
    cells = field_cells(schema)
    rock_id = cells["RockID"]
    assert rock_id in rules and any("RockIDList" in f and f"${rock_id[0]}${rock_id[1:]}" in f for f in rules[rock_id])
    optional = {"Purpose", "Goals", "Delivers"} | {f"Attention{i}" for i in (1, 2, 3)}
    for sec in schema["rock_tab"]["sections"]:
        for f in sec.get("fields", []):
            if f.get("required"):
                assert any(x.endswith('=""') for x in rules[cells[f["name"]]]), f["name"]
            elif f["name"] in optional:
                assert cells[f["name"]] not in rules, f["name"]  # optional
    ms = block(schema, "Milestones")
    first, last = ms["first_row"], ms["first_row"] + ms["capacity"] - 1
    assert any("COUNTA" in f for k, v in rules.items() for f in v if k.startswith(f"A{first}:"))  # orphan row
    planned = build_workbook.col_letter(sheets["Milestones"], "PlannedDate")
    assert any("Committed" in f for f in rules[f"{planned}{first}:{planned}{last}"])  # planned date required when committed
    team = block(schema, "Teams")
    for col in ("TeamName", "Budgeted", "Actual"):  # required team cells on a row that has a title
        letter = build_workbook.col_letter(sheets["Teams"], col)
        rng = f"{letter}{team['first_row']}:{letter}{team['first_row'] + team['capacity'] - 1}"
        if col != "TeamName":
            assert rng in rules, col
    fills = {r.rules[0].dxf.fill.bgColor.rgb[-6:] for r in ws.conditional_formatting}
    assert fills == {build_workbook.ACCENT_TINT}


def test_steerco_dropdown_is_a_warning(wb, schema, sheets):
    ws = wb[schema["rock_tab"]["template_sheet"]]
    sec = block(schema, "Decisions")
    letter = build_workbook.col_letter(sheets["Decisions"], "SteercoDate")
    rng = f"{letter}{sec['first_row']}:{letter}{sec['first_row'] + sec['capacity'] - 1}"
    dv = next(d for d in ws.data_validations.dataValidation if str(d.sqref) == rng)
    assert dv.errorStyle == "warning" and dv.formula1 == "lst_lkp_Steercos_Date"


def test_number_formats_follow_the_schema(wb, schema):
    """FillPct is 0% on the rock tab (section 2, strip row), Rocks and Teams; SlipWeeks is 0.0 by key."""
    rt = schema["rock_tab"]
    ws = wb[rt["template_sheet"]]
    fill = next(f for sec in rt["sections"] for f in sec.get("fields", []) if f["name"] == "FillPct")
    assert ws[f"{rt['value_column']}{fill['row']}"].number_format == "0%"
    rocks = next(s for s in schema["sheets"] if s["name"] == "Rocks")
    strip_col = [c["name"] for c in rocks["columns"]].index("FillPct") + 1
    assert ws.cell(row=rt["strip_row"], column=strip_col).number_format == "0%"
    expected = 0
    for s in schema["sheets"]:
        for i, c in enumerate(s["columns"], 1):
            fmt = c.get("number_format") or (build_workbook.DATE_FORMAT if c["type"] == "date" else None)
            if s["kind"] == "consolidated" and fmt:
                assert wb[s["name"]].cell(row=5, column=i).number_format == fmt, (s["name"], c["name"])
                expected += 1
    assert wb["Rocks"].cell(row=5, column=strip_col).number_format == "0%"
    assert wb["Teams"].cell(row=5, column=[c["name"] for c in next(s for s in schema["sheets"] if s["name"] == "Teams")["columns"]].index("FillPct") + 1).number_format == "0%"
    sw = [c["name"] for c in next(s for s in schema["sheets"] if s["name"] == "Milestones")["columns"]].index("SlipWeeks") + 1
    assert wb["Milestones"].cell(row=5, column=sw).number_format == "0.0"
    assert expected
    # block calculated cells carry the format too
    ms = block(schema, "Milestones")
    assert ws.cell(row=ms["first_row"], column=sw).number_format == "0.0"


def test_ado_link_validation_requires_https(wb, schema, sheets):
    ws = wb[schema["rock_tab"]["template_sheet"]]
    sec = block(schema, "Milestones")
    letter = build_workbook.col_letter(sheets["Milestones"], "ADOLink")
    rng = f"{letter}{sec['first_row']}:{letter}{sec['first_row'] + sec['capacity'] - 1}"
    dv = next(d for d in ws.data_validations.dataValidation if str(d.sqref) == rng)
    first = f"{letter}{sec['first_row']}"
    assert dv.type == "custom" and dv.formula1 == f'OR({first}="",LEFT({first},8)="https://")'
    assert dv.showErrorMessage


def test_teams_counts_are_whole_numbers(wb, schema, sheets):
    ws = wb[schema["rock_tab"]["template_sheet"]]
    sec = block(schema, "Teams")
    for col in ("Budgeted", "Actual"):
        letter = build_workbook.col_letter(sheets["Teams"], col)
        rng = f"{letter}{sec['first_row']}:{letter}{sec['first_row'] + sec['capacity'] - 1}"
        dv = next(d for d in ws.data_validations.dataValidation if str(d.sqref) == rng)
        assert dv.type == "whole" and dv.operator == "greaterThanOrEqual" and dv.formula1 == "0"


def test_section_one_formula_fields_are_grey_and_locked(wb, schema):
    rt = schema["rock_tab"]
    ws = wb[rt["template_sheet"]]
    names = ["OpenRequests"] + [f"Request{i}" for i in range(1, 6)]
    cells = field_cells(schema)
    for n in names:
        c = ws[cells[n]]
        assert str(c.value).startswith("=") and c.protection.locked and c.fill.fgColor.rgb.endswith(build_workbook.CALC_GREY), n


def test_sheet_order_matches_the_schema(wb, schema, tabs):
    rt = schema["rock_tab"]
    names = wb.sheetnames
    assert names[names.index(rt["end_sheet"]) + 1:] == [s["name"] for s in schema["sheets"]]
    assert [s["name"] for s in schema["sheets"] if s["kind"] == "consolidated"] == [
        "Rocks", "Milestones", "Risks", "CallOuts", "Decisions", "Teams"]


def test_consolidated_sheets_header_and_anchor(wb, schema):
    anchor = schema["rock_tab"]["consolidation_formula"]["anchor"]
    for s in schema["sheets"]:
        if s["kind"] != "consolidated":
            continue
        ws = wb[s["name"]]
        assert [c.value for c in ws[1]] == [c["name"] for c in s["columns"]]
        assert isinstance(ws[anchor].value, ArrayFormula)
        # nothing else is typed on the sheet
        others = [c.coordinate for row in ws.iter_rows(min_row=2) for c in row if c.value is not None and c.coordinate != anchor]
        assert not others
        assert ws.protection.sheet
        for i, c in enumerate(s["columns"], 1):
            if c["type"] == "date":
                assert ws.cell(row=3, column=i).number_format == build_workbook.DATE_FORMAT, c["name"]


def test_consolidation_formula_ranges_follow_the_schema(wb, schema):
    cf = schema["rock_tab"]["consolidation_formula"]
    rocks_cols = len(next(s for s in schema["sheets"] if s["name"] == "Rocks")["columns"])
    last = build_workbook.get_column_letter(rocks_cols)
    assert f"{cf['span']}!$A$2:${last}$2" in fixed(wb["Rocks"]["A2"])
    for sec in schema["rock_tab"]["sections"]:
        if sec["kind"] != "block":
            continue
        cols = len(next(s for s in schema["sheets"] if s["name"] == sec["block"])["columns"])
        rng = f"$A${sec['first_row']}:${build_workbook.get_column_letter(cols)}${sec['first_row'] + sec['capacity'] - 1}"
        f = fixed(wb[sec["block"]]["A2"])
        assert f"VSTACK({cf['span']}!{rng})" in f, sec["block"]
        assert "ISTEXT" in f and "_xlfn.VSTACK" in f


def test_dynamic_array_metadata_on_all_six_anchors(schema):
    path = ROOT / "data" / "workbook.xlsx"
    with zipfile.ZipFile(path) as z:
        names = load_workbook(path).sheetnames
        for sheet, cell in build_workbook.dynamic_cells(schema):
            xml = z.read(f"xl/worksheets/sheet{names.index(sheet) + 1}.xml").decode()
            assert re.search(r'<c r="%s"[^>]* cm="1"' % cell, xml), sheet
        assert "XLDAPR" in z.read("xl/metadata.xml").decode()
    assert len(build_workbook.dynamic_cells(schema)) == 6


def test_params_defined_names_use_file_prefixes(wb):
    for name in ("MCQuarterStart", "MCQuarterEnd", "OutlookEnd", "CurrentMCQuarter", "NextMCQuarter"):
        assert name in wb.defined_names
    assert wb.defined_names["NextSteerco"].attr_text.startswith("_xlfn.MINIFS(")
    assert wb.defined_names["HighRiskMin"].attr_text == "6"


def test_history_protection_matches_snapshot_script(wb, schema):
    """snapshot.ts unprotects, appends and calls protect({allowAutoFilter: true, allowSort: true}) without password."""
    assert re.search(r"protect\(\{\s*allowAutoFilter:\s*true,\s*allowSort:\s*true\s*\}\)", SNAPSHOT_TS)
    for name in ("RockSnapshots", "MilestoneSnapshots"):
        p = wb[name].protection
        assert p.sheet and not p.password
        assert p.autoFilter is False and p.sort is False  # allowed
        assert p.insertRows and p.deleteRows and p.formatCells  # everything else locked


def test_history_columns_match_snapshot_script(schema):
    rock_copy = re.search(r"const ROCK_COPY_COLUMNS = \[(.*?)\];", SNAPSHOT_TS, re.S).group(1)
    cols = ["SnapshotID", "SnapshotDate"] + re.findall(r'"(\w+)"', rock_copy)
    assert cols == [c["name"] for c in next(s for s in schema["sheets"] if s["name"] == "RockSnapshots")["columns"]]


def test_migration_is_idempotent_on_structure(tmp_path, schema):
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    build_workbook.main(["--deploy", "--from", str(ROOT / "data" / "workbook.xlsx"), "--out", str(a)])
    build_workbook.main(["--deploy", "--from", str(a), "--out", str(b)])
    assert build_workbook.read_rock_tabs(a, schema) == build_workbook.read_rock_tabs(b, schema)
    assert load_workbook(a).sheetnames == load_workbook(b).sheetnames


@pytest.mark.skipif(not shutil.which("soffice"), reason="soffice not installed")
def test_smoke_libreoffice_opens_and_about_counts_calculate(tmp_path, schema):
    src = tmp_path / "deploy.xlsx"
    build_workbook.main(["--deploy", "--out", str(src)])
    out = tmp_path / "out"
    r = subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                       capture_output=True, text=True, timeout=120)
    assert r.returncode == 0, r.stderr
    about = load_workbook(out / "deploy.xlsx", data_only=True)["About"]
    counts = {row[0]: row[3] for row in about.iter_rows(min_row=5, max_col=4, values_only=True) if row[0]}
    for s in schema["sheets"]:
        if s["kind"] in ("lookup", "history"):
            assert counts[s["name"]] == len(s.get("seed", [])), s["name"]


def _milestone_dv(wb, schema, sheets, col):
    ws = wb[schema["rock_tab"]["template_sheet"]]
    sec = block(schema, "Milestones")
    letter = build_workbook.col_letter(sheets["Milestones"], col)
    rng = f"{letter}{sec['first_row']}:{letter}{sec['first_row'] + sec['capacity'] - 1}"
    return next(d for d in ws.data_validations.dataValidation if str(d.sqref) == rng)


def test_progress_is_a_whole_number_from_0_to_100(wb, schema, sheets):
    dv = _milestone_dv(wb, schema, sheets, "Progress")
    assert dv.type == "whole" and dv.operator == "between"
    assert (dv.formula1, dv.formula2) == ("0", "100") and dv.allow_blank and dv.showErrorMessage


def test_milestone_owner_is_a_people_dropdown_that_only_warns(wb, schema, sheets):
    dv = _milestone_dv(wb, schema, sheets, "Owner")
    assert dv.type == "list" and dv.formula1 == "lst_lkp_People_Name"
    assert dv.errorStyle == "warning" and dv.allow_blank


def test_progress_columns_sit_right_after_the_lead_inputs(schema):
    names = [c["name"] for c in next(s for s in schema["sheets"] if s["name"] == "Milestones")["columns"]]
    assert names[names.index("Detail3") + 1:names.index("Detail3") + 3] == ["Progress", "ProgressShown"]
    assert names.index("ProgressShown") < names.index("Status")
