"""T11: the About sheet carries definitions, checks, the update guide and the add a rock steps."""
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import build_workbook  # noqa: E402


def texts(ws):
    return [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


@pytest.fixture(scope="module", params=["sample", "deploy"])
def about(request, tmp_path_factory):
    if request.param == "sample":
        return load_workbook(ROOT / "data" / "workbook.xlsx")["About"]
    out = tmp_path_factory.mktemp("about") / "d.xlsx"
    build_workbook.main(["--deploy", "--out", str(out)])
    return load_workbook(out)["About"]


def test_guide_steps_definitions_and_rules_present(about, schema):
    found = texts(about)
    spec = schema["about"]
    for key in ("update_guide", "add_rock", "close_rock", "rules"):
        for line in spec[key]:
            assert line in found, f"{key}: {line[:40]}"
    for d in spec["definitions"]:
        assert d["term"] in found and d["meaning"] in found


def test_checks_listed_with_formulas(about, schema):
    rows = {r[0].value: r for r in about.iter_rows() if r[0].value}
    for chk in schema["about"]["checks"]:
        row = rows[chk["label"]]
        assert row[1].value == chk["ok_when"]
        f = row[2].value
        assert isinstance(f, str) and f.startswith("="), chk["label"]
        assert "<" not in f.replace("<>", "").replace("<=", "").replace(">=", "") or "<>" in f
        assert "INDEX(" in f or "NextSteerco" in f


@pytest.mark.skipif(not shutil.which("soffice"), reason="soffice not installed")
def test_checks_calculate_without_errors_on_sample(tmp_path, schema):
    out = tmp_path / "out"
    r = subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out),
                        str(ROOT / "data" / "workbook.xlsx")], capture_output=True, text=True, timeout=180)
    assert r.returncode == 0, r.stderr
    ws = load_workbook(out / "workbook.xlsx", data_only=True)["About"]
    results = {r[0].value: r[2].value for r in ws.iter_rows() if r[0].value}
    # LibreOffice cannot evaluate the 3D VSTACK consolidation (TD01/TD08), so checks that read the
    # consolidated sheets show #NAME? here; only the standalone check is asserted on values.
    for chk in schema["about"]["checks"]:
        if "<" in chk["excel"].replace("<>", ""):
            continue
        v = results[chk["label"]]
        assert v is not None and not str(v).startswith("#") and not str(v).startswith("Err"), (chk["label"], v)
