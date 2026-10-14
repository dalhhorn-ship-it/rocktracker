"""Package level checks of the generated xlsx, against what Excel accepts without a repair prompt
(DECISIONS 0022). We have no Excel here, so these check the OOXML rules a repair prompt usually comes from:
well formed parts, element order inside <worksheet>, dynamic array metadata wiring, LET names, validation
and defined name limits, tables, and the bisect kit."""
import html
import re
import sys
import zipfile
from pathlib import Path

import pytest
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import build_workbook  # noqa: E402

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
M = "{%s}" % NS["m"]
# ECMA-376 order of the children of <worksheet> (the ones this project can produce, plus the rest for completeness)
WORKSHEET_ORDER = ["sheetPr", "dimension", "sheetViews", "sheetFormatPr", "cols", "sheetData", "sheetCalcPr",
                   "sheetProtection", "protectedRanges", "scenarios", "autoFilter", "sortState", "dataConsolidate",
                   "customSheetViews", "mergeCells", "phoneticPr", "conditionalFormatting", "dataValidations",
                   "hyperlinks", "printOptions", "pageMargins", "pageSetup", "headerFooter", "rowBreaks", "colBreaks",
                   "customProperties", "cellWatches", "ignoredErrors", "smartTags", "drawing", "legacyDrawing",
                   "legacyDrawingHF", "drawingHF", "picture", "oleObjects", "controls", "webPublishItems",
                   "tableParts", "extLst"]


@pytest.fixture(scope="module")
def files(tmp_path_factory):
    out = tmp_path_factory.mktemp("pkg") / "d.xlsx"
    build_workbook.main(["--deploy", "--out", str(out)])
    with zipfile.ZipFile(out) as z:
        return {n: z.read(n) for n in z.namelist()}


@pytest.fixture(scope="module")
def sample_files():
    with zipfile.ZipFile(ROOT / "data" / "workbook.xlsx") as z:
        return {n: z.read(n) for n in z.namelist()}


@pytest.fixture(params=["deploy", "sample"])
def pkg(request, files, sample_files):
    return files if request.param == "deploy" else sample_files


def sheets_of(pkg):
    return {n: etree.fromstring(b) for n, b in pkg.items() if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n)}


def test_every_part_is_well_formed_xml(pkg):
    for n, b in pkg.items():
        if n.endswith((".xml", ".rels", ".vml")):
            etree.fromstring(b)  # raises on malformed XML


def test_content_types_first_and_cover_every_part(pkg):
    assert list(pkg)[0] == "[Content_Types].xml"
    ct = etree.fromstring(pkg["[Content_Types].xml"])
    defaults = {e.get("Extension") for e in ct if e.tag.endswith("Default")}
    overrides = {e.get("PartName") for e in ct if e.tag.endswith("Override")}
    for n in pkg:
        if n == "[Content_Types].xml":
            continue
        assert "/" + n in overrides or n.rsplit(".", 1)[-1] in defaults, f"no content type for {n}"
    for o in overrides:
        assert o.lstrip("/") in pkg, f"content type for a missing part {o}"


def test_relationship_targets_exist(pkg):
    for n, b in pkg.items():
        if not n.endswith(".rels"):
            continue
        base = Path(n).parent.parent
        for rel in etree.fromstring(b):
            target = rel.get("Target")
            path = target.lstrip("/") if target.startswith("/") else str((base / target).as_posix())
            parts = []
            for p in path.split("/"):  # resolve ..
                if p == "..":
                    parts.pop()
                elif p:
                    parts.append(p)
            assert "/".join(parts) in pkg, f"{n}: missing target {target}"


def test_worksheet_children_are_in_schema_order(pkg):
    for n, root in sheets_of(pkg).items():
        tags = [etree.QName(c).localname for c in root]
        assert all(t in WORKSHEET_ORDER for t in tags), (n, tags)
        idx = [WORKSHEET_ORDER.index(t) for t in tags]
        assert idx == sorted(idx), f"{n}: elements out of order {tags}"
        for single in ("sheetPr", "dimension", "sheetViews", "sheetFormatPr", "sheetData", "sheetProtection",
                       "dataValidations", "pageMargins", "tableParts", "legacyDrawing"):
            assert tags.count(single) <= 1, (n, single)


def test_no_empty_value_elements_and_no_empty_inline_strings(pkg):
    for n, b in pkg.items():
        if n.startswith("xl/worksheets/sheet"):
            assert b"<v></v>" not in b and b"<v/>" not in b and b"<v />" not in b, f"{n} has an empty <v>"
            for c in etree.fromstring(b).iter(M + "c"):
                if c.get("t") == "inlineStr":
                    assert c.find(M + "is") is not None, f"{n}!{c.get('r')} inline string without content"


def test_dynamic_array_metadata_is_wired_like_excel_writes_it(files):
    ct = etree.fromstring(files["[Content_Types].xml"])
    ov = {e.get("PartName"): e.get("ContentType") for e in ct if e.tag.endswith("Override")}
    assert ov["/xl/metadata.xml"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"
    rels = etree.fromstring(files["xl/_rels/workbook.xml.rels"])
    meta = [r for r in rels if r.get("Target") == "metadata.xml"]
    assert len(meta) == 1
    assert meta[0].get("Type") == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sheetMetadata"
    md = etree.fromstring(files["xl/metadata.xml"])
    assert md.find("m:metadataTypes/m:metadataType", NS).get("name") == "XLDAPR"
    assert md.find("m:futureMetadata", NS).get("name") == "XLDAPR"
    assert md.find("m:cellMetadata", NS).get("count") == str(len(md.findall("m:cellMetadata/m:bk", NS)))
    schema = build_workbook.load_schema()
    wanted = len(build_workbook.dynamic_cells(schema))
    cm_cells = []
    for n, root in sheets_of(files).items():
        for c in root.iter(M + "c"):
            if c.get("cm") is not None:
                f = c.find(M + "f")
                assert c.get("cm") == "1" and f is not None and f.get("t") == "array" and f.get("ref") == c.get("r"), (n, c.get("r"))
                assert c.find(M + "v") is None
                cm_cells.append((n, c.get("r")))
    assert len(cm_cells) == wanted == 6


def split_args(s):
    """Top level arguments of a function call body, quote and paren aware."""
    out, depth, cur, q = [], 0, "", False
    for ch in s:
        if ch == '"':
            q = not q
        if not q:
            if ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
            elif ch == "," and depth == 0:
                out.append(cur)
                cur = ""
                continue
        cur += ch
    return out + [cur]


def let_bodies(formula):
    for m in re.finditer(r"_xlfn\.LET\(", formula):
        depth, i, q = 1, m.end(), False
        while depth:
            ch = formula[i]
            q ^= ch == '"'
            if not q:
                depth += (ch == "(") - (ch == ")")
            i += 1
        yield formula[m.end():i - 1]


def all_formulas(pkg):
    for n, b in pkg.items():
        if n.startswith("xl/worksheets/sheet"):
            for f in etree.fromstring(b).iter(M + "f"):
                yield n, f.text or ""
    for d in etree.fromstring(pkg["xl/workbook.xml"]).iter(M + "definedName"):
        yield "definedName " + d.get("name"), d.text or ""


def test_let_variable_names_are_prefixed_and_not_reserved(pkg):
    seen = 0
    for where, f in all_formulas(pkg):
        for body in let_bodies(f):
            args = split_args(body)
            for name in args[0:-1:2]:
                seen += 1
                assert name.startswith("_xlpm."), f"{where}: LET name {name!r} lacks the _xlpm. prefix"
                bare = name[len("_xlpm."):]
                assert bare.lower() not in ("c", "r") and not re.fullmatch(r"[A-Za-z]{1,3}\d+", bare), f"{where}: reserved or cell like LET name {bare}"
    assert seen >= 0


def test_unprefixed_new_functions_are_not_used(pkg):
    needs = ["LET", "VSTACK", "HSTACK", "CHOOSECOLS", "TOCOL", "REDUCE", "LAMBDA", "SEQUENCE", "FILTER", "SORT", "MINIFS", "UNIQUE", "XLOOKUP"]
    for where, f in all_formulas(pkg):
        bare = re.sub(r'"(?:[^"]|"")*"', "", f)
        for fn in needs:
            assert not re.search(r"(?<![\w.])%s\(" % fn, bare), f"{where}: {fn} without _xlfn. prefix"


def test_no_spill_reference_in_conditional_formats_or_validations(pkg):
    for n, root in sheets_of(pkg).items():
        for el in list(root.iter(M + "cfRule")) + list(root.iter(M + "dataValidation")):
            text = "".join(x.text or "" for x in el.iter() if etree.QName(x).localname in ("formula", "formula1", "formula2"))
            assert "ANCHORARRAY" not in text and "#" not in text, f"{n}: spill reference in {text}"
            assert "!" not in text, f"{n}: a sheet reference inside a rule ({text}); use a defined name"


def test_data_validation_limits_and_no_overlap(pkg):
    def cells(sqref):
        out = set()
        for part in sqref.split():
            m = re.fullmatch(r"([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?", part)
            c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3) or m.group(1), int(m.group(4) or m.group(2))
            from openpyxl.utils import column_index_from_string as ci
            for c in range(ci(c1), ci(c2) + 1):
                for r in range(r1, r2 + 1):
                    out.add((c, r))
        return out

    for n, root in sheets_of(pkg).items():
        seen = set()
        for dv in root.iter(M + "dataValidation"):
            assert len(dv.get("errorTitle", "")) <= 32 and len(dv.get("error", "")) <= 255, (n, dv.get("errorTitle"))
            assert len(dv.get("promptTitle", "")) <= 32 and len(dv.get("prompt", "")) <= 255
            for f in ("formula1", "formula2"):
                e = dv.find(M + f)
                assert e is None or len(e.text) <= 255, (n, e.text)
            c = cells(dv.get("sqref"))
            assert not (c & seen), f"{n}: overlapping data validation ranges {dv.get('sqref')}"
            seen |= c
        count = root.find(M + "dataValidations")
        if count is not None:
            assert int(count.get("count")) == len(count)


def test_defined_names_are_valid(pkg):
    names = [d.get("name") for d in etree.fromstring(pkg["xl/workbook.xml"]).iter(M + "definedName")]
    assert len(names) == len(set(names))
    for n in names:
        assert re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", n) and len(n) <= 255, n
        assert not re.fullmatch(r"[A-Za-z]{1,3}\d+", n), f"{n} looks like a cell address"
        assert not re.fullmatch(r"[Rr]\d*[Cc]\d*|[Rr]|[Cc]", n), f"{n} looks like an R1C1 reference"


def test_tables_have_unique_names_and_matching_headers(files):
    names = []
    for n, b in files.items():
        if n.startswith("xl/tables/"):
            t = etree.fromstring(b)
            names.append(t.get("displayName"))
            assert t.get("name") and len(t.get("displayName")) <= 255
            cols = [c.get("name") for c in t.iter(M + "tableColumn")]
            assert len(cols) == len(set(cols))
            assert int(t.find("m:tableColumns", NS).get("count")) == len(cols)
    assert len(names) == len(set(names)) == 13
    # header cells equal the table column names
    from openpyxl import load_workbook
    import io
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, b in files.items():
            z.writestr(n, b)
    wb = load_workbook(buf)
    for ws in wb.worksheets:
        for t in ws.tables.values():
            headers = [c.value for c in ws[t.ref.split(":")[0] + ":" + re.sub(r"\d+", "1", t.ref.split(":")[1])][0]]
            assert headers == [c.name for c in t.tableColumns], (ws.title, headers)


def test_string_and_formula_lengths(pkg):
    for where, f in all_formulas(pkg):
        assert len(f) <= 8192, where
    for n, b in pkg.items():
        if n.startswith("xl/worksheets/sheet"):
            for t in etree.fromstring(b).iter(M + "t"):
                assert len(t.text or "") <= 32767
    wb = etree.fromstring(pkg["xl/workbook.xml"])
    for s in wb.iter(M + "sheet"):
        assert len(s.get("name")) <= 31 and not re.search(r"[\\/?*\[\]:]", s.get("name"))


def test_sheet_protection_attributes_are_booleans(pkg):
    for n, root in sheets_of(pkg).items():
        for p in root.iter(M + "sheetProtection"):
            for k, v in p.attrib.items():
                if k not in ("password", "algorithmName", "hashValue", "saltValue", "spinCount"):
                    assert v in ("0", "1", "true", "false"), (n, k, v)


def test_repair_kit_variants(tmp_path):
    names = build_workbook.repair_kit(build_workbook.load_schema(), tmp_path)
    assert sorted(names) == ["1_full.xlsx", "2_no_dynamic_arrays.xlsx", "3_no_conditional_formats.xlsx",
                             "4_no_data_validation.xlsx", "5_no_defined_names.xlsx", "6_bare.xlsx", "CHECK.md"]
    assert len((tmp_path / "CHECK.md").read_text().strip().splitlines()) == 5
    got = {}
    for n in names:
        if n.endswith(".xlsx"):
            with zipfile.ZipFile(tmp_path / n) as z:
                got[n] = {x: z.read(x) for x in z.namelist()}
            for part, b in got[n].items():
                if part.endswith((".xml", ".rels", ".vml")):
                    etree.fromstring(b)

    def has(n, pat):
        return any(re.search(pat, b.decode(errors="ignore")) for part, b in got[n].items() if part.startswith("xl/worksheets/sheet") or part == "xl/workbook.xml")

    assert has("1_full.xlsx", r"cm=\"1\"") and has("1_full.xlsx", "<conditionalFormatting") and has("1_full.xlsx", "<dataValidation ") and has("1_full.xlsx", "<definedName ")
    assert not has("2_no_dynamic_arrays.xlsx", r"cm=\"1\"") and "xl/metadata.xml" not in got["2_no_dynamic_arrays.xlsx"]
    assert not has("3_no_conditional_formats.xlsx", "<conditionalFormatting") and has("3_no_conditional_formats.xlsx", "<dataValidation ")
    assert not has("4_no_data_validation.xlsx", "<dataValidation ") and has("4_no_data_validation.xlsx", "<conditionalFormatting")
    assert not has("5_no_defined_names.xlsx", "<definedName ") and not has("5_no_defined_names.xlsx", r"type=\"list\"")
    bare = "6_bare.xlsx"
    assert not any(has(bare, p) for p in (r"cm=\"1\"", "<conditionalFormatting", "<dataValidation ", "<definedName ", "<legacyDrawing"))
