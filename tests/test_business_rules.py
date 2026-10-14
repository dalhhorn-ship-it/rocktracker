"""Layer 3: business rules ("Computed rules" in docs/ARCHITECTURE.md) against the Python mirror
scripts/rock_model.py. Fixtures are hand made rows with today frozen, so the tests do not
depend on sample data."""
import datetime as dt
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import rock_model as rm  # noqa: E402

D = dt.date
TODAY = D(2026, 10, 14)  # MC26Q3: 2026-10-01..2026-12-31, outlook to 2027-03-31
STEERCOS = ["2026-09-01", "2026-10-28", "2026-12-02"]


def rock(**kw):
    base = dict(RockID="R01", RockName="Checkout", ValueChain="Buy", Lead="Alex Example",
                GPM="Sam Example", Active="Yes", Purpose="p", UpdateDate="2026-10-10",
                State="On track", MainMessage="m")
    base.update(kw)
    return base


def model(*rocks, today=TODAY, steercos=STEERCOS):
    return rm.build_model(list(rocks), today, steercos)


def ms(**kw):
    base = dict(Title="M", Commitment="Committed", PlannedDate="2026-11-30")
    base.update(kw)
    return base


def ms_of(**kw):
    """Compute one milestone through a rock and return its consolidated row."""
    return model(rock(Milestones=[ms(**kw)]))["Milestones"][0]


def ms_at(today, **kw):
    return model(rock(Milestones=[ms(**kw)]), today=today)["Milestones"][0]


def strip_of(**kw):
    return model(rock(**kw))["Rocks"][0]


# ------------------------------------------------------------------ MC quarter label
@pytest.mark.parametrize("d,label", [
    (D(2026, 9, 30), "MC26Q2"), (D(2026, 10, 1), "MC26Q3"), (D(2026, 12, 31), "MC26Q3"),
    (D(2027, 1, 1), "MC26Q4"), (D(2027, 3, 31), "MC26Q4"), (D(2027, 4, 1), "MC27Q1"),
])
def test_mc_quarter_label_at_boundaries(d, label):
    assert rm.mc_quarter(d) == label
    assert ms_of(PlannedDate=d.isoformat())["MCQuarter"] == label


def test_mc_quarter_empty_without_a_date():
    assert ms_of(PlannedDate=None)["MCQuarter"] == ""


# ------------------------------------------------------------------ params
def test_params_today_14_oct_2026():
    p = rm.params(TODAY, STEERCOS)
    assert p["MCQuarterStart"] == D(2026, 10, 1)
    assert p["MCQuarterEnd"] == D(2026, 12, 31)
    assert p["OutlookEnd"] == D(2027, 3, 31)
    assert (p["CurrentMCQuarter"], p["NextMCQuarter"]) == ("MC26Q3", "MC26Q4")
    assert p["NextSteerco"] == D(2026, 10, 28)


def test_params_across_the_year_end():
    p = rm.params(D(2027, 1, 2))
    assert (p["CurrentMCQuarter"], p["NextMCQuarter"]) == ("MC26Q4", "MC27Q1")
    assert p["MCQuarterStart"] == D(2027, 1, 1) and p["OutlookEnd"] == D(2027, 6, 30)


@pytest.mark.parametrize("today,qs,qe,oe", [
    (D(2026, 1, 1), D(2026, 1, 1), D(2026, 3, 31), D(2026, 6, 30)),
    (D(2026, 3, 31), D(2026, 1, 1), D(2026, 3, 31), D(2026, 6, 30)),
    (D(2026, 4, 1), D(2026, 4, 1), D(2026, 6, 30), D(2026, 9, 30)),
    (D(2026, 12, 31), D(2026, 10, 1), D(2026, 12, 31), D(2027, 3, 31)),
])
def test_quarter_and_outlook_on_boundaries(today, qs, qe, oe):
    p = rm.params(today)
    assert (p["MCQuarterStart"], p["MCQuarterEnd"], p["OutlookEnd"]) == (qs, qe, oe)


def test_next_steerco_today_counts_and_none_when_absent():
    assert rm.params(D(2026, 10, 28), STEERCOS)["NextSteerco"] == D(2026, 10, 28)
    assert rm.params(D(2027, 1, 1), STEERCOS)["NextSteerco"] is None


def test_rocks_carry_current_and_next_mc_quarter():
    s = strip_of()
    assert (s["CurrentMCQuarter"], s["NextMCQuarter"]) == ("MC26Q3", "MC26Q4")
    assert rm.tab_fields([rock()], TODAY)[0]["OutlookQuarters"] == "MC26Q3 and MC26Q4"


# ------------------------------------------------------------------ milestone status
def test_status_committed_on_plan_default():
    assert ms_of()["Status"] == "On plan"


def test_status_empty_forecast_is_on_plan():
    assert ms_of(ForecastDate=None)["Status"] == "On plan"


def test_status_forecast_equal_or_earlier_is_on_plan():
    assert ms_of(ForecastDate="2026-11-30")["Status"] == "On plan"
    assert ms_of(ForecastDate="2026-11-01")["Status"] == "On plan"


def test_status_slipped_and_slip_weeks():
    m = ms_of(ForecastDate="2026-12-14")
    assert m["Status"] == "Slipped" and m["SlipWeeks"] == 2.0


def test_slip_weeks_one_decimal_and_blank_when_not_slipped():
    assert ms_of(ForecastDate="2026-12-03")["SlipWeeks"] == 0.4
    assert ms_of()["SlipWeeks"] == ""


def test_status_empty_planned_never_slips():
    assert ms_of(PlannedDate=None, ForecastDate="2026-12-14")["Status"] == "On plan"


def test_status_on_watch_when_on_plan_and_slipped_beats_it():
    assert ms_of(LeadStatus="On watch")["Status"] == "On watch"
    assert ms_of(LeadStatus="On watch", ForecastDate="2026-12-14")["Status"] == "Slipped"


def test_completed_beats_on_hold_slipped_and_not_committed():
    assert ms_of(LeadStatus="Completed", ForecastDate="2026-12-14")["Status"] == "Completed"
    assert ms_of(LeadStatus="Completed", Commitment="Choice point")["Status"] == "Completed"
    assert ms_of(LeadStatus="Completed")["SlipWeeks"] == ""


def test_actual_date_alone_does_not_complete_a_milestone():
    m = ms_of(ActualDate="2026-10-01")
    assert m["Status"] == "On plan"


def test_on_hold_beats_slipped_and_not_committed():
    assert ms_of(LeadStatus="On hold", ForecastDate="2026-12-14")["Status"] == "On hold"
    assert ms_of(LeadStatus="On hold", Commitment="Choice point")["Status"] == "On hold"


@pytest.mark.parametrize("commitment", ["Choice point", "Not committed", ""])
def test_not_committed_including_empty_commitment(commitment):
    m = ms_of(Commitment=commitment, ForecastDate="2026-12-14")
    assert m["Status"] == "Not committed" and m["SlipWeeks"] == ""


def test_on_watch_ignored_when_not_committed():
    assert ms_of(Commitment="Choice point", LeadStatus="On watch")["Status"] == "Not committed"


# ------------------------------------------------------------------ timeline date
def test_timeline_date_precedence():
    assert ms_of(ActualDate="2026-10-05", ForecastDate="2026-12-01")["TimelineDate"] == D(2026, 10, 5)
    assert ms_of(ForecastDate="2026-12-01")["TimelineDate"] == D(2026, 12, 1)
    assert ms_of()["TimelineDate"] == D(2026, 11, 30)
    assert ms_of(PlannedDate=None)["TimelineDate"] == ""


def test_actual_date_sets_the_mc_quarter():
    m = ms_of(PlannedDate="2026-11-30", ActualDate="2027-01-05")
    assert m["MCQuarter"] == "MC26Q4"


# ------------------------------------------------------------------ in outlook
@pytest.mark.parametrize("planned,forecast,actual,expected", [
    ("2026-10-01", None, None, "Yes"),          # outlook start
    ("2026-09-30", None, None, "No"),
    ("2027-03-31", None, None, "Yes"),          # OutlookEnd
    ("2027-04-01", None, None, "No"),
    ("2026-09-30", "2026-10-01", None, "Yes"),  # on timeline only
    ("2027-03-31", "2027-04-01", None, "Yes"),  # on planned only
    ("2026-09-01", None, "2026-11-01", "Yes"),  # timeline is the actual date
    (None, None, None, "No"),
    (None, "2026-11-01", None, "Yes"),
])
def test_in_outlook_edges(planned, forecast, actual, expected):
    assert ms_of(PlannedDate=planned, ForecastDate=forecast, ActualDate=actual)["InOutlook"] == expected


# ------------------------------------------------------------------ visible
def test_visible_completed_with_actual_before_today_is_hidden():
    assert ms_of(LeadStatus="Completed", ActualDate="2026-10-13",
                 PlannedDate="2026-11-30")["Visible"] == "No"


def test_visible_completed_without_actual_and_planned_before_today_is_hidden():
    assert ms_of(LeadStatus="Completed", PlannedDate="2026-10-13")["Visible"] == "No"


def test_visible_completed_with_actual_today_is_shown_when_in_outlook():
    assert ms_of(LeadStatus="Completed", ActualDate="2026-10-14",
                 PlannedDate="2026-08-01")["Visible"] == "Yes"


def test_visible_completed_in_future_follows_outlook():
    assert ms_of(LeadStatus="Completed", PlannedDate="2026-12-01")["Visible"] == "Yes"
    assert ms_of(LeadStatus="Completed", PlannedDate="2027-06-01")["Visible"] == "No"


def test_visible_completed_undated_is_hidden():
    assert ms_of(LeadStatus="Completed", PlannedDate=None)["Visible"] == "No"


def test_visible_not_completed_overdue_is_shown_even_outside_outlook():
    m = ms_of(PlannedDate="2026-08-01")
    assert m["InOutlook"] == "No" and m["Visible"] == "Yes"


def test_visible_not_completed_undated_is_shown():
    assert ms_of(PlannedDate=None)["Visible"] == "Yes"


def test_visible_not_completed_beyond_outlook_is_hidden():
    assert ms_of(PlannedDate="2027-04-01")["Visible"] == "No"


def test_visible_not_completed_in_outlook_is_shown():
    assert ms_of(PlannedDate="2027-03-31")["Visible"] == "Yes"


# ------------------------------------------------------------------ overdue
def test_overdue_uses_timeline_date():
    assert ms_of(PlannedDate="2026-10-13")["Overdue"] == "Yes"
    assert ms_of(PlannedDate="2026-10-13", ForecastDate="2026-10-20")["Overdue"] == "No"
    assert ms_of(PlannedDate="2026-11-30", ForecastDate="2026-10-01")["Overdue"] == "Yes"


def test_overdue_today_is_not_overdue():
    assert ms_of(PlannedDate="2026-10-14")["Overdue"] == "No"


def test_overdue_excludes_completed_on_hold_not_committed_and_undated():
    assert ms_of(PlannedDate="2026-10-01", LeadStatus="Completed")["Overdue"] == "No"
    assert ms_of(PlannedDate="2026-10-01", LeadStatus="On hold")["Overdue"] == "No"
    assert ms_of(PlannedDate="2026-10-01", Commitment="Not committed")["Overdue"] == "No"
    assert ms_of(PlannedDate=None)["Overdue"] == "No"


def test_overdue_counts_slipped_and_on_watch():
    assert ms_of(PlannedDate="2026-09-01", ForecastDate="2026-10-01")["Overdue"] == "Yes"
    assert ms_of(PlannedDate="2026-09-01", LeadStatus="On watch")["Overdue"] == "Yes"


# ------------------------------------------------------------------ risks
@pytest.mark.parametrize("lk,im,score,high", [
    ("Low", "Low", 1, "No"), ("Medium", "Medium", 4, "No"), ("High", "Low", 3, "No"),
    ("Medium", "High", 6, "Yes"), ("High", "Medium", 6, "Yes"), ("High", "High", 9, "Yes"),
])
def test_risk_score_and_high(lk, im, score, high):
    r = model(rock(Risks=[dict(Title="x", Likelihood=lk, Impact=im)]))["Risks"][0]
    assert (r["Score"], r["High"]) == (score, high)


def test_risk_without_impact_has_no_score_and_is_not_high():
    r = model(rock(Risks=[dict(Title="x", Likelihood="High")]))["Risks"][0]
    assert r["Score"] == "" and r["High"] == "No" and r["NeedsAction"] == "No"


def test_needs_action_when_high_without_action_or_owner():
    base = dict(Title="x", Likelihood="High", Impact="High")
    rows = model(rock(Risks=[base, dict(base, KeyAction="a"), dict(base, ActionOwner="o"),
                             dict(base, KeyAction="a", ActionOwner="o")]))["Risks"]
    assert [r["NeedsAction"] for r in rows] == ["Yes", "Yes", "Yes", "No"]


def test_closed_high_risk_needs_no_action_and_is_not_open():
    risks = [dict(Title="x", Likelihood="High", Impact="High", RiskStatus="Closed")]
    m = model(rock(Risks=risks))
    assert m["Risks"][0]["NeedsAction"] == "No"
    assert m["Rocks"][0]["OpenHighRisks"] == 0 and m["Rocks"][0]["HighRisksNoAction"] == 0


def test_empty_risk_status_counts_as_open():
    risks = [dict(Title="x", Likelihood="High", Impact="High", KeyAction="a", ActionOwner="o")]
    s = model(rock(Risks=risks))["Rocks"][0]
    assert s["OpenHighRisks"] == 1 and s["HighRisksNoAction"] == 0


def test_low_risk_does_not_need_action():
    r = model(rock(Risks=[dict(Title="x", Likelihood="Low", Impact="Low")]))["Risks"][0]
    assert r["NeedsAction"] == "No"


def test_risk_has_no_top_risk_column():
    assert "TopRisk" not in model(rock(Risks=[dict(Title="x")]))["Risks"][0]


# ------------------------------------------------------------------ milestone counts
def test_milestone_counts_x_of_y():
    items = [
        ms(Title="a"),                                      # on plan
        ms(Title="b", LeadStatus="On watch"),               # on watch
        ms(Title="c", ForecastDate="2026-12-20"),           # slipped
        ms(Title="d", LeadStatus="Completed"),              # completed
        ms(Title="e", Commitment="Choice point"),           # not committed, not counted
        ms(Title="f", LeadStatus="On hold"),                # committed, but not on plan
        ms(Title="g", PlannedDate="2028-01-01"),            # outside the outlook
    ]
    s = strip_of(Milestones=items)
    assert s["MilestonesCommitted"] == 5     # a b c d f
    assert s["MilestonesSlipped"] == 1
    assert s["MilestonesOnPlan"] == 3        # a b d
    assert s["MilestonesCompleted"] == 1


def test_on_plan_counts_completed_so_ratio_is_stable():
    items = [ms(Title="a"), ms(Title="b")]
    before = strip_of(Milestones=items)
    items[1]["LeadStatus"] = "Completed"
    after = strip_of(Milestones=items)
    assert (before["MilestonesOnPlan"], before["MilestonesCommitted"]) == (2, 2)
    assert (after["MilestonesOnPlan"], after["MilestonesCommitted"]) == (2, 2)


def test_completed_in_outlook_stays_counted_after_it_is_hidden():
    items = [ms(Title="a", LeadStatus="Completed", ActualDate="2026-10-05")]
    m = model(rock(Milestones=items))
    assert m["Milestones"][0]["Visible"] == "No"
    assert m["Rocks"][0]["MilestonesCompleted"] == 1


def test_overdue_count_covers_any_quarter_and_excludes_on_hold():
    items = [ms(Title="a", PlannedDate="2025-01-01"),
             ms(Title="b", PlannedDate="2025-01-01", LeadStatus="On hold"),
             ms(Title="c", PlannedDate="2026-11-30")]
    s = strip_of(Milestones=items)
    assert s["MilestonesOverdue"] == 1
    assert s["MilestonesCommitted"] == 1  # only c is in the outlook


def test_empty_rock_has_real_zero_counts():
    s = strip_of()
    for c in ("MilestonesCommitted", "MilestonesSlipped", "MilestonesOnPlan",
              "MilestonesCompleted", "MilestonesOverdue", "OpenHighRisks", "HighRisksNoAction",
              "PositionsBudgeted", "PositionsActual"):
        assert s[c] == 0 and s[c] is not ""
    assert s["FillPct"] == ""


# ------------------------------------------------------------------ teams and positions
def teams_of(*ts, **kw):
    return model(rock(Teams=list(ts), **kw))


def test_team_fill_pct_budgeted_zero_and_empty():
    t = teams_of(dict(TeamName="a", Budgeted=4, Actual=3),
                 dict(TeamName="b", Budgeted=0, Actual=2),
                 dict(TeamName="c", Actual=2),
                 dict(TeamName="d", Budgeted=5, Actual=5))["Teams"]
    assert [x["FillPct"] for x in t] == [0.75, "", "", 1.0]


def test_team_with_zero_actual_keeps_zero_in_the_mirror():
    t = teams_of(dict(TeamName="a", Budgeted=4, Actual=0))["Teams"][0]
    assert t["Actual"] == 0 and t["FillPct"] == 0 and t["Budgeted"] == 4  # DECISIONS 0018


def test_team_empty_count_arrives_as_zero_and_other_zero_columns_still_blank():
    m = teams_of(dict(TeamName="a", Actual=2))
    assert m["Teams"][0]["Budgeted"] == 0
    assert rm.zero_wrap({"X": 0, "Budgeted": 0}, keep={"Budgeted"}) == {"X": "", "Budgeted": 0}


def test_teams_anchor_formula_has_the_column_guard(schema):
    import build_workbook as bw
    teams = next(s for s in schema["sheets"] if s["name"] == "Teams")
    assert teams["keep_zero_columns"] == ["Budgeted", "Actual", "FillPct"]
    cols = [c["name"] for c in teams["columns"]]
    f = bw.consolidation_formula(schema, teams, bw.Resolver(schema))
    guard = "".join(f"*(c<>{cols.index(k) + 1})" for k in teams["keep_zero_columns"])
    assert f'IF((f=0){guard},"",f)' in f and "SEQUENCE(1,COLUMNS(f))" in f
    others = next(s for s in schema["sheets"] if s["name"] == "Risks")
    g = bw.consolidation_formula(schema, others, bw.Resolver(schema))
    assert 'IF(f=0,"",f)' in g and "SEQUENCE" not in g


def test_rock_totals_sum_keyed_teams_and_keep_real_zero():
    s = teams_of(dict(TeamName="a", Budgeted=4, Actual=3),
                 dict(TeamName="b", Budgeted=2, Actual=0))["Rocks"][0]
    assert (s["PositionsBudgeted"], s["PositionsActual"], s["FillPct"]) == (6, 3, 0.5)
    s = teams_of(dict(TeamName="a", Budgeted=0, Actual=0))["Rocks"][0]
    assert (s["PositionsBudgeted"], s["PositionsActual"], s["FillPct"]) == (0, 0, "")


def test_rock_totals_ignore_rows_without_a_team_name():
    s = teams_of(dict(TeamName="a", Budgeted=4, Actual=3),
                 dict(Budgeted=10, Actual=10))["Rocks"][0]
    assert (s["PositionsBudgeted"], s["PositionsActual"]) == (4, 3)


def test_team_keys_use_slot_numbers():
    t = teams_of(dict(TeamName="a"), {}, dict(TeamName="c"))["Teams"]
    assert [x["TeamKey"] for x in t] == ["R01-T01", "R01-T03"]


def test_inactive_rock_keeps_decisions_drops_teams():
    m = model(rock(Active="No", Teams=[dict(TeamName="a", Budgeted=4, Actual=3)],
                   Decisions=[dict(Decision="d", Outcome="Go")]))
    assert m["Teams"] == []
    assert [d["DecisionKey"] for d in m["Decisions"]] == ["R01-D01"]
    assert m["Rocks"][0]["PositionsBudgeted"] == 0


# ------------------------------------------------------------------ update age, stale
def test_update_age_and_stale_edges():
    assert strip_of(UpdateDate="2026-10-14")["UpdateAgeDays"] == 0
    s = strip_of(UpdateDate="2026-09-30")  # exactly 14 days
    assert s["UpdateAgeDays"] == 14 and s["Stale"] == "No"
    s = strip_of(UpdateDate="2026-09-29")  # 15 days
    assert s["UpdateAgeDays"] == 15 and s["Stale"] == "Yes"


def test_empty_update_date_is_stale_with_empty_age():
    s = strip_of(UpdateDate=None)
    assert s["UpdateAgeDays"] == "" and s["Stale"] == "Yes"


def test_attention_points_are_carried_on_rocks():
    s = strip_of(Attention1="a", Attention3="c")
    assert (s["Attention1"], s["Attention2"], s["Attention3"]) == ("a", "", "c")


# ------------------------------------------------------------------ data signal, check the call
def test_signal_on_track_for_empty_rock():
    s = strip_of()
    assert s["DataSignal"] == "On track" and s["CheckCall"] == "No"


def test_signal_at_risk_when_slipped_inside_quarter():
    s = strip_of(Milestones=[ms(PlannedDate="2026-11-30", ForecastDate="2026-12-31")])
    assert s["DataSignal"] == "At risk"


def test_signal_at_risk_on_an_overdue_milestone_alone():
    s = strip_of(Milestones=[ms(PlannedDate="2026-10-01")])
    assert s["MilestonesSlipped"] == 0 and s["HighRisksNoAction"] == 0
    assert s["MilestonesOverdue"] == 1 and s["DataSignal"] == "At risk"


def test_signal_on_hold_overdue_stays_on_track():
    s = strip_of(Milestones=[ms(PlannedDate="2026-10-01", LeadStatus="On hold")])
    assert s["DataSignal"] == "On track"


def test_signal_off_track_when_forecast_after_quarter_end():
    s = strip_of(Milestones=[ms(PlannedDate="2026-11-30", ForecastDate="2027-01-01")])
    assert s["DataSignal"] == "Off track"


def test_signal_off_track_needs_planned_in_current_quarter():
    s = strip_of(Milestones=[ms(PlannedDate="2027-01-15", ForecastDate="2027-03-01")])
    assert s["DataSignal"] == "At risk"


def test_signal_off_track_quarter_edges():
    for planned, expected in (("2026-12-31", "Off track"), ("2026-10-01", "Off track"),
                              ("2026-09-30", "At risk")):
        s = strip_of(Milestones=[ms(PlannedDate=planned, ForecastDate="2027-01-01")])
        assert s["DataSignal"] == expected, planned


def test_signal_at_risk_from_high_risk_without_action():
    s = strip_of(Risks=[dict(Title="x", Likelihood="High", Impact="High")])
    assert s["DataSignal"] == "At risk"
    s = strip_of(Risks=[dict(Title="x", Likelihood="High", Impact="High",
                             KeyAction="a", ActionOwner="o")])
    assert s["DataSignal"] == "On track"


def test_signal_ignores_completed_slip_and_uncommitted():
    items = [ms(LeadStatus="Completed", ForecastDate="2027-06-01"),
             ms(Commitment="Choice point", ForecastDate="2027-06-01")]
    assert strip_of(Milestones=items)["DataSignal"] == "On track"


@pytest.mark.parametrize("state,expected", [("On track", "Yes"), ("At risk", "Yes"),
                                            ("Off track", "No")])
def test_check_call_when_signal_off_track(state, expected):
    s = strip_of(State=state, Milestones=[ms(PlannedDate="2026-11-30", ForecastDate="2027-02-01")])
    assert s["DataSignal"] == "Off track" and s["CheckCall"] == expected


def test_check_call_data_better_than_state_is_no():
    assert strip_of(State="Off track")["CheckCall"] == "No"


def test_check_call_empty_without_state():
    assert strip_of(State=None)["CheckCall"] == ""


# ------------------------------------------------------------------ decisions
def decisions(*ds, **rockkw):
    return model(rock(Decisions=list(ds), **rockkw))["Decisions"]


def test_decision_state_requested_until_outcome():
    a, b = decisions(dict(Decision="d1"), dict(Decision="d2", Outcome="Go"))
    assert (a["DecisionState"], b["DecisionState"]) == ("Requested", "Decided")


def test_request_no_skips_decided_rows_and_counts_by_slot():
    rows = decisions(dict(Decision="a"), dict(Decision="b", Outcome="Go"), dict(Decision="c"),
                     {}, dict(Decision="e"))
    assert [(r["Decision"], r["RequestNo"]) for r in rows] == [
        ("a", 1), ("b", ""), ("c", 2), ("e", 3)]


def test_request_no_restarts_per_rock():
    m = model(rock(Decisions=[dict(Decision="a"), dict(Decision="b")]),
              rock(RockID="R02", Decisions=[dict(Decision="c")]))
    assert [r["RequestNo"] for r in m["Decisions"]] == [1, 2, 1]


def test_request_lines_show_first_five_of_six_open():
    items = [dict(Decision=f"d{i}") for i in range(1, 7)]
    items.insert(2, dict(Decision="done", Outcome="Go"))
    f = rm.tab_fields([rock(Decisions=items)], TODAY)[0]
    assert f["OpenRequests"] == 6
    assert [f[f"Request{i}"] for i in range(1, 6)] == ["d1", "d2", "d3", "d4", "d5"]
    rows = model(rock(Decisions=items))["Decisions"]
    assert [r["RequestNo"] for r in rows if r["DecisionState"] == "Requested"] == [1, 2, 3, 4, 5, 6]


def test_request_lines_empty_when_none_open():
    f = rm.tab_fields([rock(Decisions=[dict(Decision="a", Outcome="Go")])], TODAY)[0]
    assert f["OpenRequests"] == 0 and f["Request1"] == "" and f["Request5"] == ""


def test_due_next_steerco_by_steerco_date_or_decide_by():
    rows = decisions(
        dict(Decision="a", SteercoDate="2026-10-28"),                         # raised for next
        dict(Decision="b", SteercoDate="2026-12-02", DecideBy="2026-10-28"),  # due by next
        dict(Decision="c", SteercoDate="2026-12-02", DecideBy="2026-10-27"),  # before next
        dict(Decision="d", SteercoDate="2026-12-02", DecideBy="2026-10-29"),  # after next
        dict(Decision="e", SteercoDate="2026-12-02"),                         # later steerco
        dict(Decision="f", SteercoDate="2026-10-28", Outcome="Go"),           # decided
    )
    assert [r["DueNextSteerco"] for r in rows] == ["Yes", "Yes", "Yes", "No", "No", "No"]


def test_due_next_steerco_no_future_steerco():
    rows = rm.build_model([rock(Decisions=[dict(Decision="a", SteercoDate="2026-10-28",
                                                 DecideBy="2026-01-01")])],
                          TODAY, steercos=["2026-01-01"])["Decisions"]
    assert rows[0]["DueNextSteerco"] == "No"


def test_due_next_steerco_empty_dates_not_due():
    assert decisions(dict(Decision="a"))[0]["DueNextSteerco"] == "No"


def test_open_follow_up_rules():
    rows = decisions(
        dict(Decision="a"),
        dict(Decision="b", Outcome="Go"),
        dict(Decision="c", Outcome="Go", FollowUp="Open"),
        dict(Decision="d", Outcome="Go", FollowUp="Closed"),
    )
    assert [r["OpenFollowUp"] for r in rows] == ["No", "Yes", "Yes", "No"]


def test_requested_decision_with_closed_follow_up_is_not_open_follow_up():
    assert decisions(dict(Decision="a", FollowUp="Closed"))[0]["OpenFollowUp"] == "No"


# ------------------------------------------------------------------ consolidation filter
def test_keys_use_slot_numbers_including_gaps():
    m = model(rock(Milestones=[ms(Title="a"), {}, ms(Title="c")]))
    assert [x["MilestoneKey"] for x in m["Milestones"]] == ["R01-M01", "R01-M03"]


def test_row_without_title_is_dropped_everywhere():
    m = model(rock(
        Milestones=[ms(Title="")], Risks=[dict(Likelihood="High", Impact="High")],
        CallOuts=[dict(Date="2026-11-01")], Decisions=[dict(Outcome="Go")],
        Teams=[dict(Budgeted=3)]))
    assert all(m[k] == [] for k in ("Milestones", "Risks", "CallOuts", "Decisions", "Teams"))


def test_inactive_rock_keeps_decisions_only():
    m = model(rock(
        Active="No", Milestones=[ms()],
        Risks=[dict(Title="r", Likelihood="High", Impact="High")],
        CallOuts=[dict(Text="c")], Decisions=[dict(Decision="d", Outcome="Go")]))
    assert m["Milestones"] == [] and m["Risks"] == [] and m["CallOuts"] == []
    assert [d["DecisionKey"] for d in m["Decisions"]] == ["R01-D01"]
    assert m["Decisions"][0]["Rock"] == "Checkout"


def test_empty_active_counts_as_inactive():
    assert model(rock(Active=None, Milestones=[ms()]))["Milestones"] == []


def test_rock_without_id_contributes_nothing():
    m = model(rock(RockID=None, Milestones=[ms()], Decisions=[dict(Decision="d")],
                   Teams=[dict(TeamName="t")]))
    assert all(v == [] for v in m.values())


def test_rock_template_style_empty_tab_contributes_nothing():
    assert all(v == [] for v in model({}).values())


def test_inactive_rock_stays_in_rocks_with_active_no():
    s = strip_of(Active="No", Milestones=[ms(ForecastDate="2027-01-01")])
    assert s["Active"] == "No"
    assert s["MilestonesCommitted"] == 0 and s["DataSignal"] == "On track"


def test_rows_copy_rock_profile_and_follow_tab_order():
    m = model(rock(Milestones=[ms(Title="a")]),
              rock(RockID="R02", RockName="Search", GPM="Jo Example", Milestones=[ms(Title="b")]))
    assert [(x["MilestoneKey"], x["Rock"], x["GPM"]) for x in m["Milestones"]] == [
        ("R01-M01", "Checkout", "Sam Example"), ("R02-M01", "Search", "Jo Example")]
    assert [r["RockID"] for r in m["Rocks"]] == ["R01", "R02"]


def test_blank_cells_come_out_empty_not_none_in_text_columns():
    m = ms_of()
    assert m["Comment"] == "" and m["LeadStatus"] == "" and m["Detail1"] == ""


# ------------------------------------------------------------------ schema alignment
def _full_rock():
    return rock(Milestones=[ms()], Risks=[dict(Title="r")], CallOuts=[dict(Text="c")],
                Decisions=[dict(Decision="d")], Teams=[dict(TeamName="t", Budgeted=2, Actual=1)])


def test_model_columns_match_schema_order():
    schema = rm.load_schema()
    m = rm.build_model([_full_rock()], TODAY, STEERCOS, schema=schema)
    consolidated = [s for s in schema["sheets"] if s["name"] in m]
    assert {s["name"] for s in consolidated} == {"Rocks", "Milestones", "Risks", "CallOuts",
                                                 "Decisions", "Teams"}
    for s in consolidated:
        assert m[s["name"]], s["name"]
        assert list(m[s["name"]][0]) == [c["name"] for c in s["columns"]]


def test_every_schema_column_is_mirrored_without_padding():
    """The mirror itself produces every consolidated column (order_columns only reorders)."""
    schema = rm.load_schema()
    m = rm.build_model([_full_rock()], TODAY, STEERCOS)
    for s in schema["sheets"]:
        if s["name"] in m:
            produced = set(m[s["name"]][0])
            missing = [c["name"] for c in s["columns"] if c["name"] not in produced]
            assert not missing, f"{s['name']} not mirrored: {missing}"
            extra = produced - {c["name"] for c in s["columns"]}
            assert not extra, f"{s['name']} has columns outside the schema: {extra}"


def test_every_rock_tab_formula_field_is_mirrored():
    schema = rm.load_schema()
    f = rm.tab_fields([_full_rock()], TODAY, STEERCOS)[0]
    for sec in schema["rock_tab"]["sections"]:
        if sec["kind"] == "fields":
            for fld in sec["fields"]:
                if fld.get("formula"):
                    assert fld["name"] in f, fld["name"]


# ------------------------------------------------------------------ sample data (tolerant until T41)
def test_sample_rocks_compute_without_errors_and_match_schema():
    schema = rm.load_schema()
    rocks = rm.rocks_from_json(Path(__file__).resolve().parent.parent / "data" / "sample.json")
    m = rm.build_model(rocks, TODAY, STEERCOS, schema=schema)
    assert m["Rocks"]
    for s in schema["sheets"]:
        if s["name"] in m and m[s["name"]]:
            assert list(m[s["name"]][0]) == [c["name"] for c in s["columns"]]


# ------------------------------------------------------------------ workbook loader
def test_loader_reads_rock_tab_by_schema_layout(tmp_path):
    from openpyxl import Workbook
    schema = rm.load_schema()
    rt = schema["rock_tab"]
    wb = Workbook()
    wb.active.title = rt["template_sheet"]
    ws = wb.create_sheet("R01 Checkout")
    wb.create_sheet(rt["end_sheet"])
    sheets = {s["name"]: s for s in schema["sheets"]}
    vals = {"RockID": "R01", "RockName": "Checkout", "Active": "Yes", "State": "At risk",
            "UpdateDate": dt.datetime(2026, 10, 10), "Attention1": "watch this"}
    block_vals = {"Milestones": {"Title": "M1", "Commitment": "Committed",
                                 "PlannedDate": dt.datetime(2026, 11, 30)},
                  "Teams": {"TeamName": "Core", "Budgeted": 4, "Actual": 0}}
    for sec in rt["sections"]:
        if sec["kind"] == "fields":
            for f in sec["fields"]:
                if f["name"] in vals:
                    ws[f"{rt['value_column']}{f['row']}"] = vals[f["name"]]
        else:
            typed = [c for c in sheets[sec["block"]]["columns"] if not c.get("formula")]
            for i, c in enumerate(typed, 2):
                ws.cell(row=sec["column_header_row"], column=i, value=c["label"])
            for i, c in enumerate(typed, 2):
                v = block_vals.get(sec["block"], {}).get(c["name"])
                ws.cell(row=sec["first_row"], column=i, value=v)
    path = tmp_path / "t.xlsx"
    wb.save(path)
    rocks = rm.rocks_from_workbook(path, schema)
    assert len(rocks) == 1 and len(rocks[0]["Milestones"]) == 15 and len(rocks[0]["Teams"]) == 5
    m = rm.build_model(rocks, TODAY, STEERCOS)
    assert m["Milestones"][0]["MilestoneKey"] == "R01-M01"
    assert m["Milestones"][0]["Status"] == "On plan"
    assert m["Rocks"][0]["Attention1"] == "watch this"
    assert (m["Rocks"][0]["PositionsBudgeted"], m["Rocks"][0]["PositionsActual"]) == (4, 0)


# ------------------------------------------------------------------ milestone progress (PRD item 7)
def test_progress_shown_is_the_typed_value_when_not_completed():
    assert ms_of(Progress=60)["ProgressShown"] == 60


def test_progress_shown_is_blank_when_progress_is_empty():
    assert ms_of()["ProgressShown"] == ""
    assert ms_of(Progress=None)["ProgressShown"] == ""


def test_progress_shown_is_100_when_completed_whatever_was_typed():
    assert ms_of(LeadStatus="Completed", Progress=80, ActualDate="2026-10-14")["ProgressShown"] == 100
    assert ms_of(LeadStatus="Completed")["ProgressShown"] == 100


def test_a_typed_zero_progress_stays_zero_on_the_consolidated_sheet():
    assert ms_of(Progress=0)["ProgressShown"] == 0
    assert ms_of(Progress=0)["Progress"] == ""  # the input column is wrapped; the portal reads ProgressShown


def test_progress_shown_is_blank_for_a_row_without_a_key():
    assert model(rock(Milestones=[ms(Title="", Progress=50)]))["Milestones"] == []
