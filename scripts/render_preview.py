"""Render portal/preview.html: a reference render of portal/portal.spec.json over the
workbook, styled with design/tokens.css. The Copilot Cowork portal should look and read like this.

Run: python3 scripts/render_preview.py
"""
import datetime as dt
import html
import json
import math
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
import portal_metrics as pm  # noqa: E402

CATEGORICAL = [f"var(--cat-{i})" for i in range(1, 9)]


def load_rows(schema):
    wb = load_workbook(ROOT / schema["workbook"], data_only=True)
    data = {}
    for sheet in schema["sheets"]:
        values = list(wb[sheet["name"]].iter_rows(values_only=True))
        headers = list(values[0])
        data[sheet["name"]] = [dict(zip(headers, r)) for r in values[1:] if any(v is not None for v in r)]
    return data


def esc(value):
    return html.escape(pm.fmt_date(value))


def fig(value):
    return f'<em class="fig">{esc(value)}</em>'


def bar_chart(chart, data, stage_map):
    points = pm.series(chart, data)
    if chart["color"] == "stage":
        order = ["backlog", "progress", "review", "done"]
        points.sort(key=lambda kv: order.index(stage_map.get(kv[0], "backlog")))
    top = max((n for _, n in points), default=1)
    rows = []
    for label, n in points:
        colour = f"var(--stage-{stage_map.get(label, 'backlog')})" if chart["color"] == "stage" else "var(--cat-1)"
        rows.append(
            f'<div class="brow"><span>{esc(label)}</span>'
            f'<div class="track"><div class="bar" title="{n}" style="width:{100 * n / top:.0f}%;background:{colour}"></div></div></div>'
        )
    return "".join(rows)


def donut_chart(chart, data):
    points = pm.series(chart, data)
    total = sum(n for _, n in points) or 1
    r, c = 40, 2 * math.pi * 40
    arcs, legend, offset = [], [], 0.0
    for i, (label, n) in enumerate(points):
        colour = "var(--cat-other)" if label == "Other" else CATEGORICAL[i % len(CATEGORICAL)]
        length = c * n / total
        arcs.append(
            f'<circle r="{r}" cx="50" cy="50" fill="none" stroke="{colour}" stroke-width="18" '
            f'stroke-dasharray="{length:.2f} {c - length:.2f}" stroke-dashoffset="{-offset:.2f}"><title>{esc(label)}: {n}</title></circle>'
        )
        legend.append(f'<div><i style="background:{colour}"></i>{esc(label)}</div>')
        offset += length
    return (
        '<div class="donut"><svg viewBox="0 0 100 100" width="140" height="140" style="transform:rotate(-90deg)">'
        + "".join(arcs) + '</svg><div class="legend">' + "".join(legend) + "</div></div>"
    )


def section(tab, n, data, values, stage_map):
    parts = [
        f'<section class="panel" id="tab{n}"{"" if n == 0 else " hidden"}>',
        f'<div class="eyebrow" style="margin-top:24px">{html.escape(tab["eyebrow"])}</div>',
        f'<h2 class="h1">{html.escape(tab["headline"])}</h2>',
        f'<p class="lede">{html.escape(tab["lede"])}</p>',
    ]
    if tab.get("narrative"):
        parts.append(f'<p class="narrative">{pm.fill(html.escape(tab["narrative"]), values, fig)}</p>')
    if tab.get("insights"):
        cards = []
        for ins in tab["insights"]:
            if "list" in ins:
                items = pm.list_items(ins["list"], data)
                body = "<ul>" + "".join(
                    f"<li>{esc(item)}" + (f" — <b>{esc(date)}</b>" if date else "") + "</li>" for item, date in items
                ) + "</ul>" if items else f'<p>{html.escape(ins.get("empty", ""))}</p>'
            else:
                body = f'<p>{pm.fill(html.escape(ins["text"]), values, fig)}</p>'
            cards.append(f'<div class="card"><h3>{html.escape(ins["question"])}</h3>{body}</div>')
        parts.append('<div class="grid g3">' + "".join(cards) + "</div>")
    if tab.get("kpis"):
        tiles = [
            f'<div class="card kpi{" spot" if k.get("spotlight") else ""}"><div class="eyebrow">{html.escape(k["label"])}</div>'
            f'<div class="num">{values[k["id"]]}</div><div class="sub">{pm.fill(html.escape(k["sub"]), values)}</div></div>'
            for k in tab["kpis"]
        ]
        parts.append('<div class="grid g6">' + "".join(tiles) + "</div>")
    if tab.get("charts"):
        cards = []
        for ch in tab["charts"]:
            body = bar_chart(ch, data, stage_map) if ch["type"] == "bar" else donut_chart(ch, data)
            cards.append(f'<div class="card chart"><div class="eyebrow">{html.escape(ch["eyebrow"])}</div><h4>{html.escape(ch["title"])}</h4>{body}</div>')
        parts.append('<div class="grid g2">' + "".join(cards) + "</div>")
    if tab.get("table"):
        t = tab["table"]
        head = "".join(f"<th>{html.escape(c)}</th>" for c in t["columns"])
        body = []
        for r in data[t["table"]]:
            cells = []
            for c in t["columns"]:
                if c == t.get("status_column"):
                    stage = stage_map.get(r.get(c), "backlog")
                    cells.append(f'<td><i class="dot" style="background:var(--stage-{stage})"></i>{esc(r.get(c))}</td>')
                else:
                    cells.append(f"<td>{esc(r.get(c))}</td>")
            body.append("<tr>" + "".join(cells) + "</tr>")
        parts.append(f'<div class="card"><table class="list"><thead><tr>{head}</tr></thead><tbody>{"".join(body)}</tbody></table></div>')
    parts.append("</section>")
    return "".join(parts)


EXTRA_CSS = """
.meta b{text-transform:none;letter-spacing:0}
.brow{display:grid;grid-template-columns:110px 1fr;gap:12px;align-items:center;margin:8px 0;font:500 10px/1 var(--font-sans);color:var(--ink-3)}
.brow span{text-align:right}.track{background:linear-gradient(90deg,var(--grid) 1px,transparent 1px) 0 0/25% 100%}
.donut{display:flex;gap:24px;align-items:center}.legend div{font-size:10px;color:var(--ink-3);margin:4px 0}
.card ul{margin:0;padding-left:16px}.list{width:100%;border-collapse:collapse;font-size:11px}
.list th{font:500 9.5px/1 var(--font-sans);letter-spacing:.12em;text-transform:uppercase;color:var(--subtle);text-align:left;padding:8px}
.list td{padding:8px;border-top:1px solid var(--rule)}.dot{display:inline-block;width:7px;height:7px;border-radius:50%;margin-right:6px}
@media (max-width:900px){.g6{grid-template-columns:repeat(3,1fr)}.g3,.g2{grid-template-columns:1fr}.mast{flex-wrap:wrap;gap:12px}.tabs{overflow-x:auto}}
"""


def main():
    schema = json.loads((ROOT / "schema" / "workbook.schema.json").read_text())
    spec = json.loads((ROOT / "portal" / "portal.spec.json").read_text())
    data = load_rows(schema)
    values = pm.all_kpis(spec, data)
    m = spec["masthead"]
    tabs = "".join(
        f'<span class="tab{" active" if i == 0 else ""}" data-tab="{i}">{html.escape(t["question"])}</span>'
        for i, t in enumerate(spec["tabs"])
    )
    sections = "".join(section(t, i, data, values, spec["stage_map"]) for i, t in enumerate(spec["tabs"]))
    page = f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(m["title"])}</title><style>{(ROOT / spec["tokens"]).read_text()}{EXTRA_CSS}</style></head><body>
<header class="wrap mast"><div style="display:flex;gap:16px;align-items:center">
<svg width="20" height="20" aria-hidden="true"><circle cx="5" cy="5" r="3" fill="#D4494C"/><circle cx="13" cy="9" r="2.5" fill="#D4494C"/><circle cx="8" cy="15" r="2.5" fill="#D4494C"/></svg>
<div><div class="eyebrow">{" · ".join(html.escape(b) for b in m["breadcrumb"])}</div><h1>{html.escape(m["title"])}</h1></div></div>
<div style="display:flex"><div class="meta eyebrow">As of<b>{pm.fmt_date(dt.date.today())}</b></div><div class="meta eyebrow">Board<b>{html.escape(m["board"])}</b></div></div></header>
<nav class="wrap tabs">{tabs}</nav><main class="wrap">{sections}</main>
<script>document.querySelectorAll('.tab').forEach(t=>t.onclick=()=>{{document.querySelectorAll('.tab').forEach(x=>x.classList.toggle('active',x===t));document.querySelectorAll('.panel').forEach((p,i)=>p.hidden=String(i)!==t.dataset.tab);}});</script>
</body></html>"""
    out = ROOT / "portal" / "preview.html"
    out.write_text(page)
    print(f"wrote {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
