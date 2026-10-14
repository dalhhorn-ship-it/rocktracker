#!/usr/bin/env python3
"""Python mirror of the Rock Tracker calculated columns (schema `excel` plus `rule`).

Input: rock tabs, one dict per rock (see `normalize_rock`). Output: the five
consolidated sheets (Rocks, Epics, Risks, CallOuts, Decisions) as lists of dicts in
schema column order, with the consolidation filter applied. Used by the tests, the
LibreOffice recalc check and the preview, so the preview never reads formula cells.

Rock dict shape (all keys optional, missing means empty):
  profile/update fields by name: RockID, RockName, ValueChain, Lead, GPM, Active,
    Purpose, Goals, Delivers, UpdateDate, State, MainMessage, PositionsPlanned,
    PositionsFilled
  block rows: "Epics", "Risks", "CallOuts", "Decisions": lists of dicts keyed by the
    typed column names, one item per slot in order. {} or None is an empty slot.
Dates may be date, datetime or YYYY-MM-DD text.

Blank safety follows the workbook: calculated cells are "" when the row has no key.
"""
from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SCHEMA_PATH = ROOT / "schema" / "workbook.schema.json"

HIGH_RISK_MIN = 6
STALE_DAYS = 14
LEVEL_SCORE = {"Low": 1, "Medium": 2, "High": 3}
STATE_ORDER = ["On track", "At risk", "Off track"]

BLOCKS = ("Epics", "Risks", "CallOuts", "Decisions")
PROFILE_FIELDS = ["RockID", "RockName", "ValueChain", "Lead", "GPM", "Active", "Purpose",
                  "Goals", "Delivers"]
UPDATE_FIELDS = ["UpdateDate", "State", "MainMessage", "PositionsPlanned", "PositionsFilled"]
KEY_PREFIX = {"Epics": "E", "Risks": "R", "CallOuts": "C", "Decisions": "D"}
KEY_TITLE = {"Epics": "Title", "Risks": "Title", "CallOuts": "Text", "Decisions": "Decision"}
DATE_COLUMNS = {"UpdateDate", "PlannedDate", "ForecastDate", "Date", "SteercoDate", "DecideBy"}


# ---------------------------------------------------------------- helpers
def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return dt.date.fromisoformat(str(v)[:10])


def n(d):
    """Excel N() of a date: 0 when empty, else a comparable number."""
    return 0 if d is None else d.toordinal()


def txt(v):
    return "" if v is None else str(v).strip() if isinstance(v, str) else v


def edate(d: dt.date, months: int) -> dt.date:
    m = d.month - 1 + months
    y, m = d.year + m // 12, m % 12 + 1
    return dt.date(y, m, min(d.day, [31, 29 if y % 4 == 0 and (y % 100 or y % 400 == 0) else 28,
                                     31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1]))


def params(today: dt.date, steercos=()):
    qs = dt.date(today.year, (today.month - 1) // 3 * 3 + 1, 1)
    dates = sorted(d for d in (to_date(s) for s in steercos) if d)
    nxt = next((d for d in dates if d >= today), None)
    return {
        "QuarterStart": qs,
        "QuarterEnd": edate(qs, 3) - dt.timedelta(days=1),
        "WindowStart": qs,
        "WindowEnd": edate(qs, 6) - dt.timedelta(days=1),
        "NextSteerco": nxt,  # None means 0 in Excel
        "HighRiskMin": HIGH_RISK_MIN,
        "StaleDays": STALE_DAYS,
    }


def normalize_rock(rock: dict) -> dict:
    out = {}
    for f in PROFILE_FIELDS + UPDATE_FIELDS:
        v = rock.get(f)
        out[f] = to_date(v) if f == "UpdateDate" else ("" if v is None else v)
    for b in BLOCKS:
        rows = []
        for r in rock.get(b) or []:
            r = dict(r or {})
            for c in list(r):
                if c in DATE_COLUMNS:
                    r[c] = to_date(r[c])
                elif r[c] is None:
                    r[c] = ""
            rows.append(r)
        out[b] = rows
    return out


# ---------------------------------------------------------------- row rules
def epic_status(e):
    lead, commit = e.get("LeadStatus", ""), e.get("Commitment", "")
    planned, forecast = e.get("PlannedDate"), e.get("ForecastDate")
    if lead == "Delivered":
        return "Delivered"
    if lead == "On hold":
        return "On hold"
    if commit != "Committed":
        return "Not committed"
    if n(planned) > 0 and n(forecast) > n(planned):
        return "Slipped"
    if lead == "On watch":
        return "On watch"
    return "On plan"


def epic_row(e, rock, slot, p, today):
    active = rock["Active"] == "Yes"
    key = ""
    if e.get("Title", "") != "" and rock["RockID"] != "" and active:
        key = f"{rock['RockID']}-E{slot:02d}"
    row = {"EpicKey": key}
    for c in ("Title", "ADOId", "Commitment", "PlannedDate", "ForecastDate", "LeadStatus",
              "Comment", "KRLink"):
        row[c] = e.get(c, None if c in DATE_COLUMNS else "")
    if key == "":
        row.update(Status="", SlipWeeks="", InWindow="", Overdue="", RockID="", Rock="", GPM="")
        return row
    status = epic_status(e)
    pl, fc = e.get("PlannedDate"), e.get("ForecastDate")
    row["Status"] = status
    row["SlipWeeks"] = round((n(fc) - n(pl)) / 7 + 1e-9, 1) if status == "Slipped" else ""
    in_win = any(p["WindowStart"] <= d <= p["WindowEnd"] for d in (pl, fc) if d)
    row["InWindow"] = "Yes" if in_win else "No"
    eff = n(fc) if n(fc) > 0 else n(pl)
    overdue = (e.get("Commitment") == "Committed" and status not in ("Delivered", "On hold")
               and max(n(pl), n(fc)) > 0 and eff < n(today))
    row["Overdue"] = "Yes" if overdue else "No"
    row.update(RockID=rock["RockID"], Rock=rock["RockName"], GPM=rock["GPM"])
    return row


def risk_row(r, rock, slot, p):
    key = ""
    if r.get("Title", "") != "" and rock["RockID"] != "" and rock["Active"] == "Yes":
        key = f"{rock['RockID']}-R{slot:02d}"
    row = {"RiskKey": key}
    for c in ("Title", "Likelihood", "Impact", "KeyAction", "ActionOwner", "RiskStatus", "TopRisk"):
        row[c] = r.get(c, "")
    if key == "":
        row.update(Score="", High="", NeedsAction="", RockID="", Rock="", GPM="")
        return row
    lk, im = r.get("Likelihood", ""), r.get("Impact", "")
    score = LEVEL_SCORE[lk] * LEVEL_SCORE[im] if lk != "" and im != "" else ""
    row["Score"] = score
    high = (score if score != "" else 0) >= p["HighRiskMin"]
    row["High"] = "Yes" if high else "No"
    needs = (high and row["RiskStatus"] != "Closed"
             and (row["KeyAction"] == "" or row["ActionOwner"] == ""))
    row["NeedsAction"] = "Yes" if needs else "No"
    row.update(RockID=rock["RockID"], Rock=rock["RockName"], GPM=rock["GPM"])
    return row


def callout_row(c, rock, slot):
    key = ""
    if c.get("Text", "") != "" and rock["RockID"] != "" and rock["Active"] == "Yes":
        key = f"{rock['RockID']}-C{slot:02d}"
    row = {"CallOutKey": key, "Text": c.get("Text", ""), "Date": c.get("Date"),
           "CallOutStatus": c.get("CallOutStatus", "")}
    on = key != ""
    row.update(RockID=rock["RockID"] if on else "", Rock=rock["RockName"] if on else "",
               GPM=rock["GPM"] if on else "")
    return row


def decision_row(d, rock, slot, p):
    # Kept after the rock closes: no Active condition.
    key = ""
    if d.get("Decision", "") != "" and rock["RockID"] != "":
        key = f"{rock['RockID']}-D{slot:02d}"
    row = {"DecisionKey": key}
    for c in ("Decision", "SteercoDate", "DecideBy", "ChoicePoint", "Options", "Outcome",
              "Owner", "FollowUp"):
        row[c] = d.get(c, None if c in DATE_COLUMNS else "")
    if key == "":
        row.update(DecisionState="", DueNextSteerco="", OpenFollowUp="", RockID="", Rock="", GPM="")
        return row
    state = "Requested" if row["Outcome"] == "" else "Decided"
    nxt = n(p["NextSteerco"])
    due = (state == "Requested" and nxt > 0
           and (n(row["SteercoDate"]) == nxt
                or (n(row["DecideBy"]) > 0 and n(row["DecideBy"]) <= nxt)))
    row["DecisionState"] = state
    row["DueNextSteerco"] = "Yes" if due else "No"
    row["OpenFollowUp"] = "Yes" if state == "Decided" and row["FollowUp"] != "Closed" else "No"
    row.update(RockID=rock["RockID"], Rock=rock["RockName"], GPM=rock["GPM"])
    return row


def rock_strip(rock, epics, risks, p, today):
    """Rocks row. `epics` and `risks` are this rock's computed block rows (all slots)."""
    if rock["RockID"] == "":
        return None
    ud = rock["UpdateDate"]
    age = (today - ud).days if ud else ""
    committed = [e for e in epics if e["EpicKey"] and e["Commitment"] == "Committed"
                 and e["InWindow"] == "Yes"]
    slipped = [e for e in epics if e["EpicKey"] and e["Status"] == "Slipped"
               and e["InWindow"] == "Yes"]
    onplan = [e for e in committed if e["Status"] in ("On plan", "On watch", "Delivered")]
    delivered = [e for e in committed if e["Status"] == "Delivered"]
    open_high = [r for r in risks if r["RiskKey"] and r["High"] == "Yes"
                 and r["RiskStatus"] != "Closed"]
    no_action = [r for r in risks if r["RiskKey"] and r["NeedsAction"] == "Yes"]
    off = any(e["Status"] == "Slipped" and p["QuarterStart"] <= e["PlannedDate"] <= p["QuarterEnd"]
              and n(e["ForecastDate"]) > n(p["QuarterEnd"])
              for e in epics if e["EpicKey"] and e["PlannedDate"])
    signal = "Off track" if off else ("At risk" if len(slipped) + len(no_action) > 0 else "On track")
    state = rock["State"]
    check = ""
    if state != "" and signal != "":
        check = "Yes" if STATE_ORDER.index(signal) > STATE_ORDER.index(state) else "No"

    def blank(v):  # strip returns "" for blanks, keeps real zeros
        return "" if v in ("", None) else v

    row = {c: blank(rock[c]) for c in PROFILE_FIELDS}
    row.update({c: blank(rock[c]) for c in UPDATE_FIELDS})
    row.update(
        UpdateAgeDays=age,
        Stale="Yes" if (ud is None or (age if age != "" else 0) > p["StaleDays"]) else "No",
        EpicsCommitted=len(committed), EpicsSlipped=len(slipped), EpicsOnPlan=len(onplan),
        EpicsDelivered=len(delivered), OpenHighRisks=len(open_high),
        HighRisksNoAction=len(no_action), DataSignal=signal, CheckCall=check,
    )
    return row


# ---------------------------------------------------------------- consolidation
def build_model(rocks, today: dt.date, steercos=(), schema=None):
    """Return {"Rocks": [...], "Epics": [...], "Risks": [...], "CallOuts": [...],
    "Decisions": [...]} for the given rock tabs (in tab order) and frozen today."""
    p = params(today, steercos)
    out = {"Rocks": [], "Epics": [], "Risks": [], "CallOuts": [], "Decisions": []}
    for raw in rocks:
        rock = normalize_rock(raw)
        epics = [epic_row(e, rock, i, p, today) for i, e in enumerate(rock["Epics"], 1)]
        risks = [risk_row(r, rock, i, p) for i, r in enumerate(rock["Risks"], 1)]
        calls = [callout_row(c, rock, i) for i, c in enumerate(rock["CallOuts"], 1)]
        decs = [decision_row(d, rock, i, p) for i, d in enumerate(rock["Decisions"], 1)]
        strip = rock_strip(rock, epics, risks, p, today)
        if strip is not None:
            out["Rocks"].append(strip)
        # consolidation filter: key must be non empty text
        out["Epics"] += [r for r in epics if r["EpicKey"]]
        out["Risks"] += [r for r in risks if r["RiskKey"]]
        out["CallOuts"] += [r for r in calls if r["CallOutKey"]]
        out["Decisions"] += [r for r in decs if r["DecisionKey"]]
    if schema is not None:
        out = order_columns(out, schema)
    return out


def order_columns(model, schema):
    cols = {s["name"]: [c["name"] for c in s["columns"]] for s in schema["sheets"]}
    return {sheet: [{c: r.get(c, "") for c in cols[sheet]} for r in rows]
            for sheet, rows in model.items()}


# ---------------------------------------------------------------- loaders
def load_schema():
    return json.loads(SCHEMA_PATH.read_text())


def rocks_from_json(path):
    """Read rock dicts from a JSON file: a list, or {"rocks": [...]} / {"Rocks": [...]}."""
    data = json.loads(Path(path).read_text())
    if isinstance(data, dict):
        data = data.get("rocks", data.get("Rocks"))
    return data


def rocks_from_workbook(path, schema=None):
    """Read the typed cells of every rock tab (between RockTemplate and RocksEnd) using the
    schema layout: fields by row in the value column, block columns by header label."""
    from openpyxl import load_workbook

    schema = schema or load_schema()
    rt = schema["rock_tab"]
    wb = load_workbook(path, data_only=True)
    names = wb.sheetnames
    first, last = names.index(rt["template_sheet"]), names.index(rt["end_sheet"])
    sheets = {s["name"]: s for s in schema["sheets"]}
    vcol = rt["value_column"]
    rocks = []
    for name in names[first + 1:last]:
        ws = wb[name]
        rock = {}
        for sec in rt["sections"]:
            if sec["kind"] == "fields":
                for f in sec["fields"]:
                    if not f.get("formula"):
                        rock[f["name"]] = ws[f"{vcol}{f['row']}"].value
            else:
                typed = [c for c in sheets[sec["block"]]["columns"] if not c.get("formula")]
                headers = {}
                for cell in ws[sec["column_header_row"]]:
                    if cell.value is not None:
                        headers[str(cell.value).strip()] = cell.column
                rows = []
                for r in range(sec["first_row"], sec["first_row"] + sec["capacity"]):
                    row = {}
                    for c in typed:
                        col = headers.get(c["label"]) or headers.get(c["name"])
                        row[c["name"]] = ws.cell(row=r, column=col).value if col else None
                    rows.append(row if any(v not in (None, "") for v in row.values()) else {})
                rock[sec["block"]] = rows
        rocks.append(rock)
    return rocks


if __name__ == "__main__":
    import sys

    src = sys.argv[1]
    rocks = rocks_from_workbook(src) if src.endswith(".xlsx") else rocks_from_json(src)
    model = build_model(rocks, dt.date.today(), schema=load_schema())
    print(json.dumps(model, default=str, indent=1))
