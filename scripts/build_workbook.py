"""Build the Excel backend from schema/workbook.schema.json.

  python3 scripts/build_workbook.py                 sample workbook (data/workbook.xlsx), for tests and preview
  python3 scripts/build_workbook.py --deploy        deployment workbook (dist/...), lookups seeded, no sample rows
  python3 scripts/build_workbook.py --deploy --from live.xlsx
                                                    migration: new structure, rows copied from a downloaded
                                                    live copy by column name (columns that no longer exist
                                                    are reported, not silently lost)

The deployment workbook gets: one Excel Table per sheet, dropdowns for lookups and
references, date validation, a highlight on empty required cells, frozen headers,
locked looking header style, and an About sheet. Same builder for every project.
"""
import argparse
import datetime as dt
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
MAX_ROWS = 2000  # validation and highlight range per sheet; raise in the schema with "max_rows"
INK, SUBTLE, RULE, ACCENT_TINT = "0F1116", "8C9097", "E6E7E9", "F6E8EE"  # design/tokens.css


def load_schema():
    return json.loads((ROOT / "schema" / "workbook.schema.json").read_text())


def coerce(value, col_type):
    if col_type == "date" and isinstance(value, str) and value:
        return dt.date.fromisoformat(value)
    return value


def read_existing(path):
    """Rows per sheet from an existing workbook, as dicts keyed by header."""
    wb = load_workbook(path, data_only=False)
    out = {}
    for ws in wb.worksheets:
        values = list(ws.iter_rows(values_only=True))
        if not values:
            continue
        headers = list(values[0])
        out[ws.title] = [dict(zip(headers, r)) for r in values[1:] if any(v is not None for v in r)]
    return out


def list_name(ref):
    sheet, column = ref.split(".", 1)
    return f"lst_{sheet}_{column}"


def build(schema, rows_by_sheet, deploy):
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {s["name"]: s for s in schema["sheets"]}
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=INK)

    if deploy:
        about = wb.create_sheet("About")
        about["A1"] = schema.get("title", "Workbook")
        about["A1"].font = Font(bold=True, size=16, color=INK)
        about["A2"] = f"Version {schema.get('version', '0')} · built {dt.date.today():%d %b %Y}"
        about["A2"].font = Font(color=SUBTLE)
        about.append([])
        about.append(["Sheet", "Purpose", "Owner", "Rows"])
        for c in about[4]:
            c.font, c.fill = header_font, header_fill
        rules = [
            "Rules for everyone who edits this workbook:",
            "Add and edit rows inside the tables only. Use the dropdowns; do not type over them.",
            "Never rename, reorder or delete sheets or columns: the portal depends on them.",
            "Never reuse or change an ID. Mark records as done instead of deleting them.",
            "Cells highlighted in pink are required and still empty.",
            "Structure changes go through the factory repository (schema change, rebuild, migrate).",
        ]
        for s in schema["sheets"]:
            pk = get_column_letter([c["name"] for c in s["columns"]].index(s["primary_key"]) + 1)
            about.append([s["name"], s.get("purpose", ""), s.get("owner", ""), f"=COUNTA('{s['name']}'!{pk}:{pk})-1"])
        about.append([])
        for line in rules:
            about.append([line])
        about.column_dimensions["A"].width = 22
        about.column_dimensions["B"].width = 60
        about.column_dimensions["C"].width = 24

    for sheet in schema["sheets"]:
        ws = wb.create_sheet(sheet["name"])
        cols = sheet["columns"]
        ws.append([c["name"] for c in cols])
        for row in rows_by_sheet.get(sheet["name"], []):
            if isinstance(row, dict):
                row = [row.get(c["name"]) for c in cols]
            ws.append([coerce(v, c["type"]) for v, c in zip(row, cols)])
        last_row = max(ws.max_row, 2)
        last_col = get_column_letter(len(cols))
        table = Table(displayName=f"tbl_{sheet['name']}", ref=f"A1:{last_col}{last_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(table)
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols, start=1):
            letter = get_column_letter(i)
            ws.column_dimensions[letter].width = max(12, min(40, len(c["name"]) + 6))
            cell = ws.cell(row=1, column=i)
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(vertical="center")
            if c["type"] == "date":
                for r in range(2, last_row + 1):
                    ws.cell(row=r, column=i).number_format = "d mmm yyyy"

        if not deploy:
            continue
        max_rows = sheet.get("max_rows", MAX_ROWS)
        pk_letter = get_column_letter([c["name"] for c in cols].index(sheet["primary_key"]) + 1)
        for i, c in enumerate(cols, start=1):
            letter = get_column_letter(i)
            rng = f"{letter}2:{letter}{max_rows}"
            ref = c.get("lookup") or c.get("references")
            if ref:
                dv = DataValidation(type="list", formula1=f"={list_name(ref)}", allow_blank=not c.get("required"))
                dv.error, dv.errorTitle = f"Pick a value from {ref}", "Not in the list"
            elif c["type"] == "date":
                dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True)
                dv.error, dv.errorTitle = "Enter a date, for example 16 Oct 2026", "Not a date"
                for r in range(last_row + 1, max_rows + 1):
                    ws.cell(row=r, column=i).number_format = "d mmm yyyy"
            else:
                dv = None
            if dv:
                dv.showErrorMessage = True
                ws.add_data_validation(dv)
                dv.add(rng)
            if c.get("required") and not c.get("formula") and c["name"] != sheet["primary_key"]:
                ws.conditional_formatting.add(
                    rng,
                    FormulaRule(formula=[f'AND(${pk_letter}2<>"",{letter}2="")'],
                                fill=PatternFill("solid", fgColor=ACCENT_TINT)),
                )

    if deploy:
        targets = {c.get("lookup") or c.get("references") for s in schema["sheets"] for c in s["columns"]} - {None}
        for ref in sorted(targets):
            sheet, column = ref.split(".", 1)
            idx = [c["name"] for c in sheets[sheet]["columns"]].index(column) + 1
            letter = get_column_letter(idx)
            max_rows = sheets[sheet].get("max_rows", MAX_ROWS)
            dn = DefinedName(list_name(ref), attr_text=f"'{sheet}'!${letter}$2:${letter}${max_rows}")
            wb.defined_names[list_name(ref)] = dn
    return wb


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--deploy", action="store_true", help="build the deployment workbook")
    ap.add_argument("--from", dest="source", help="migrate rows from this existing workbook")
    ap.add_argument("--out", help="output path (default from schema)")
    args = ap.parse_args(argv)
    schema = load_schema()

    if args.source:
        existing = read_existing(args.source)
        rows, dropped = {}, []
        for s in schema["sheets"]:
            old = existing.get(s["name"], [])
            names = {c["name"] for c in s["columns"]}
            dropped += [f"{s['name']}.{h}" for h in (old[0].keys() if old else []) if h not in names]
            rows[s["name"]] = old or [list(r) for r in s.get("seed", [])]
        if dropped:
            print("WARNING: columns in the source that the schema no longer has (data not copied):")
            print("\n".join(f"  {d}" for d in dropped))
    else:
        sample = {} if args.deploy else json.loads((ROOT / "data" / "sample.json").read_text())
        rows = {s["name"]: s.get("seed") or sample.get(s["name"], []) for s in schema["sheets"]}

    out = Path(args.out) if args.out else ROOT / (schema["deploy_file"] if args.deploy else schema["workbook"])
    out.parent.mkdir(parents=True, exist_ok=True)
    build(schema, rows, deploy=args.deploy).save(out)
    print(f"wrote {out.relative_to(ROOT) if out.is_relative_to(ROOT) else out}")


if __name__ == "__main__":
    sys.exit(main())
