"""Build the T05 spike workbook: dist/spike/consolidation_spike.xlsx.

Tests the consolidation formulas of DECISIONS 0002 in Excel for the web:
  Epics     plan A: the schema formula (3D VSTACK over RockTemplate:RocksEnd) in A2
  Epics_C2  plan C2: REDUCE + INDIRECT over the tab list on sheet RockTabs
Rock tabs here are tiny (epics in A5:D9) but use the same key filter as the real ones.

  python3 scripts/build_spike.py
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_workbook import mark_dynamic_arrays, set_dynamic_formula  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "dist" / "spike" / "consolidation_spike.xlsx"
RANGE = "A5:D9"
HEADER = ["Key", "Title", "Status", "Rock"]
ROCK_EPICS = {
    "R01": [("R01-E01", "Order intake", "On track", "R01"), ("R01-E02", "Pricing tool", "At risk", "R01")],
    "R02": [("R02-E01", "Supplier portal", "Done", "R02"), ("R02-E02", "Quality gates", "On track", "R02")],
}


def plan_a_formula(schema):
    cf = schema["rock_tab"]["consolidation_formula"]
    return cf["rocks"].replace("{span}", cf["span"]).replace("{range}", RANGE)


def plan_c2_formula():
    return (
        f'=LET(t,TOCOL(RockTabs!A2:A20,1),d,REDUCE(INDIRECT("\'"&INDEX(t,1)&"\'!{RANGE}"),SEQUENCE(ROWS(t)-1,1,2),'
        f'LAMBDA(a,i,VSTACK(a,INDIRECT("\'"&INDEX(t,i)&"\'!{RANGE}")))),'
        f'k,CHOOSECOLS(d,1),FILTER(d,ISTEXT(k)*(k<>""),""))'
    )


def build(out=OUT):
    schema = json.loads((ROOT / "schema" / "workbook.schema.json").read_text())
    rt = schema["rock_tab"]
    wb = Workbook()
    wb.active.title = "Instructions"
    ins = wb["Instructions"]
    for i, line in enumerate([
        "Consolidation spike (T05). Throwaway test file, no real data.",
        "Epics: plan A, one 3D VSTACK formula in A2. Epics_C2: plan C2, REDUCE and INDIRECT over the tab list on RockTabs.",
        "Test: duplicate RockTemplate, rename it, type any text in cell A5 of the copy. The row must appear on Epics.",
    ], 1):
        ins.cell(i, 1, line)
    ins["A1"].font = Font(bold=True)

    def rock_tab(name, epics):
        ws = wb.create_sheet(name)
        ws["A1"] = f"Rock tab {name}"
        ws["A2"] = "Epics are in A5:D9. Type a key in column A to add one."
        for c, h in enumerate(HEADER, 1):
            ws.cell(4, c, h)
        for r, row in enumerate(epics, 5):
            for c, v in enumerate(row, 1):
                ws.cell(r, c, v)
        return ws

    rock_tab(rt["template_sheet"], [])
    for name, epics in ROCK_EPICS.items():
        rock_tab(name, epics)
    end = wb.create_sheet(rt["end_sheet"])
    end["A1"] = rt["end_sheet_note"]

    dyn = []
    for sheet, formula, names in (
        ("Epics", plan_a_formula(schema), "dk"),
        ("Epics_C2", plan_c2_formula(), "tdkai"),
    ):
        ws = wb.create_sheet(sheet)
        for c, h in enumerate(HEADER, 1):
            ws.cell(1, c, h).font = Font(bold=True)
        set_dynamic_formula(ws, "A2", formula, names)
        dyn.append((sheet, "A2"))
    tabs = wb.create_sheet("RockTabs")
    tabs["A1"] = "Tab"
    for r, name in enumerate(ROCK_EPICS, 2):
        tabs.cell(r, 1, name)
    wb.calculation.fullCalcOnLoad = True
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    mark_dynamic_arrays(out, dyn)
    return out


if __name__ == "__main__":
    p = build()
    print(f"wrote {p.relative_to(ROOT)}")
